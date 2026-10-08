"""Independent saved ARM workbook/CSV readback; no exporter imports/rasterization.

Reconstruct endpoint matrices from original run records, regional length matrices
from the hash-bound sparse export, and animal/group means by direct arithmetic.
"""
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path

import nibabel as nib
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
BASE = Path(__file__).resolve().parent
OUT = BASE / "combined_arm_projection_tables_462"
RECEIPT = BASE / "independent_saved_arm_tables_readback_20261009.json"
SOURCE_FIELDS = ("ARMLevel", "ARMIndex", "ARMAbbreviation", "ARMFullName", "Hemisphere", "SourceARMStatus")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_csv_key(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def integer(value):
    number = float(value)
    assert np.isfinite(number) and number.is_integer()
    return int(number)


def uid(row):
    return row["SampleID"] + "::" + row["NeuronID"]


def numeric_close(actual, expected):
    if np.isnan(expected):
        assert actual is None
    else:
        assert actual is not None and np.isclose(float(actual), expected, rtol=1e-10, atol=1e-10)


def main():
    provenance_path = OUT / "export_provenance.json"
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    assert provenance["status"] == "software_verified_descriptive_tables"
    for name, expected in provenance["artifacts"].items():
        assert sha(OUT / name) == expected
    for path, expected in provenance["source_bindings"].items():
        assert sha(path) == expected
    summary = read_csv(OUT / "neuron_summary.csv")
    manifest = read_csv(provenance["manifest"])
    targets = read_csv(OUT / "targets.csv")
    sparse = read_csv(OUT / "per_neuron_regional_measures.csv")
    assert len(summary) == 462 and len(manifest) == 462
    uids = [uid(row) for row in summary]
    assert len(set(uids)) == 462 and uids == [uid(row) for row in manifest]
    positions = {value: i for i, value in enumerate(uids)}
    target_positions = {row["TargetID"]: i for i, row in enumerate(targets)}
    assert len(target_positions) == len(targets)
    target_levels = np.array([integer(row["Level"]) for row in targets])
    main_metadata = json.loads((ROOT / "group_analysis/evolution_20261008/arm_mapping_20261009/main/endpoints/run_provenance.json").read_text(encoding="utf-8"))
    atlas_image = nib.load(main_metadata["inputs"]["atlas"]["path"])
    reference_image = nib.load(main_metadata["inputs"]["reference"]["path"])
    np.testing.assert_array_equal(atlas_image.affine, reference_image.affine)
    assert atlas_image.shape == (*reference_image.shape, 1, 6)
    assert atlas_image.header.get_xyzt_units()[0] == "mm"
    voxel_volume = abs(float(np.linalg.det(reference_image.affine[:3, :3])))
    key = {integer(row["Index"]): row for row in read_csv_key(main_metadata["inputs"]["atlas_key"]["path"])}
    atlas_data = np.asanyarray(atlas_image.dataobj)
    observed_targets = set()
    for level in range(1, 7):
        ids, footprints = np.unique(atlas_data[:, :, :, 0, level - 1], return_counts=True)
        for index, footprint in zip(ids, footprints):
            index = int(index)
            record = key.get(index)
            status = ("zero_unassigned" if not index else "unmapped_label" if record is None else
                      "mapped" if integer(record["First_Level"]) <= level <= integer(record["Last_Level"]) else "key_level_conflict")
            token = f"L{level}:{status}:{index}"
            observed_targets.add(token)
            row = targets[target_positions[token]]
            assert integer(row["ARMIndex"]) == index and row["TargetStatus"] == status
            assert np.isclose(float(row["ActualVolumeMm3"]), footprint * voxel_volume, rtol=1e-12)
            if record:
                assert row["OfficialFullName"] == record["Full_Name"] and row["Abbreviation"] == record["Abbreviation"]
                assert row["Hemisphere"] == record["Abbreviation"][1]
            assert row["KeyRangeConflict"] == str(status == "key_level_conflict")
        outside = f"L{level}:out_of_FOV:outside"
        observed_targets.add(outside)
        assert targets[target_positions[outside]]["ARMIndex"] == ""
        assert targets[target_positions[outside]]["ActualVolumeMm3"] == ""
    assert observed_targets == set(target_positions)
    del atlas_data
    eligible = np.array([row["EndpointEligible"] == "True" for row in summary])
    assert int(eligible.sum()) == 429
    for saved, source in zip(summary, manifest):
        for field, value in source.items():
            assert saved[field] == value
    counts = np.zeros((462, len(targets)))
    counts[~eligible] = np.nan
    lengths = np.zeros_like(counts)
    original_runs = []
    for endpoint_dir, axon_dir in (
        (ROOT / "group_analysis/evolution_20261008/arm_mapping_20261009/main/endpoints",
         ROOT / "group_analysis/evolution_20261008/arm_mapping_20261009/main/axons"),
        (ROOT / "group_analysis/evolution_20261008/endpoint_maps/additional_candidates_20261009",
         ROOT / "group_analysis/evolution_20261008/projection_maps/additional_candidates_20261009"),
    ):
        run_manifest = read_csv(endpoint_dir / "input_manifest.csv")
        original_runs.append((endpoint_dir, axon_dir, run_manifest))
        qc = {uid(row): row for row in read_csv(endpoint_dir / "per_neuron_qc.csv")}
        for row in run_manifest:
            i = positions[uid(row)]
            assert eligible[i] == bool(integer(qc[uid(row)]["n_computable"]))
        for row in read_csv(endpoint_dir / "per_neuron_target_counts.csv"):
            i = positions[uid(row)]
            index = str(integer(row["target_index"])) if row["target_index"] else "outside"
            token = f"L{integer(row['level'])}:{row['target_status']}:{index}"
            j = target_positions[token]
            assert counts[i, j] == 0
            counts[i, j] = integer(row["endpoint_count"])
    seen = set()
    for row in sparse:
        i, j = positions[row["NeuronUID"]], target_positions[row["TargetID"]]
        assert (i, j) not in seen
        seen.add((i, j))
        assert uid(row) == row["NeuronUID"]
        assert row["AnimalID"] == summary[i]["AnimalID"] and row["Subregion"] == summary[i]["Subregion"]
        length = float(row["AxonTemplateLengthMm"])
        assert np.isfinite(length) and length >= 0
        lengths[i, j] = length
        if eligible[i]:
            assert float(row["CandidateEndpointCount"]) == counts[i, j]
            assert integer(row["EndpointPresence"]) == int(counts[i, j] > 0)
        else:
            assert row["CandidateEndpointCount"] == "" and row["EndpointPresence"] == ""
    positive_count_cells = np.argwhere(counts > 0)
    assert all((int(i), int(j)) in seen for i, j in positive_count_cells)
    for i, row in enumerate(summary):
        for level in range(1, 7):
            level_mask = target_levels == level
            expected_total = float(row["Endpoint_candidate_axon_endpoint_count"])
            if eligible[i]:
                assert counts[i, level_mask].sum() == expected_total
            else:
                assert np.isnan(counts[i, level_mask]).all()
            assert np.isclose(lengths[i, level_mask].sum(), float(row["Axon_selected_axon_length_mm"]), rtol=1e-9, atol=1e-9)
    animal_groups = {}
    for i, row in enumerate(summary):
        animal_groups.setdefault((row["AnimalID"], row["Subregion"]), []).append(i)
    expected_animal = {}
    for key, indices in animal_groups.items():
        indices = np.array(indices)
        valid = indices[eligible[indices]]
        expected_animal[key] = (len(indices), len(valid),
            counts[valid].mean(axis=0) if len(valid) else np.full(len(targets), np.nan),
            (counts[valid] > 0).mean(axis=0) if len(valid) else np.full(len(targets), np.nan),
            lengths[indices].mean(axis=0))
    expected_group = {}
    for source in {row["Subregion"] for row in summary}:
        contributing = [value for (animal, group), value in expected_animal.items() if group == source]
        available = [value for value in contributing if value[1] > 0]
        expected_group[source] = (
            sum(value[0] for value in contributing), sum(value[1] for value in contributing),
            np.mean([value[2] for value in available], axis=0) if available else np.full(len(targets), np.nan),
            np.mean([value[3] for value in available], axis=0) if available else np.full(len(targets), np.nan),
            np.mean([value[4] for value in contributing], axis=0), len(available), len(contributing))
    workbook_path = OUT / "arm_projection_hierarchy_tables.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    for sheet_name, expected_rows in (("Summary", summary), ("Targets", targets)):
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(iterator)
        n_rows = 0
        for row_index, values in enumerate(iterator):
            assert row_index < len(expected_rows)
            assert len(values) <= len(headers)
            values = values + (None,) * (len(headers) - len(values))
            expected_row = expected_rows[row_index]
            for field, value in zip(headers, values):
                expected = expected_row[field]
                if isinstance(value, (float, int)) and not isinstance(value, bool):
                    assert np.isclose(float(value), float(expected), rtol=1e-10, atol=1e-10)
                else:
                    assert ("" if value is None else str(value)) == expected
            n_rows += 1
        assert n_rows == len(expected_rows)
    matrix_cells = 0
    for level in range(1, 7):
        for suffix, expected in (("EP_Count", counts), ("EP_Presence", np.where(np.isnan(counts), np.nan, counts > 0)), ("AxonLen_mm", lengths)):
            sheet = workbook[f"L{level}_{suffix}"]
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator)
            fields = {name: i for i, name in enumerate(headers)}
            for field in SOURCE_FIELDS:
                assert field in fields
            column_indices = [(i, target_positions[name.split(" | ", 1)[0]]) for i, name in enumerate(headers) if name.startswith(f"L{level}:")]
            assert len(column_indices) == int((target_levels == level).sum())
            n_rows = 0
            for values in iterator:
                assert len(values) <= len(headers)
                values = values + (None,) * (len(headers) - len(values))
                i = positions[values[fields["NeuronUID"]]]
                for field in SOURCE_FIELDS:
                    saved = values[fields[field]]
                    assert ("" if saved is None else str(saved)) == summary[i][field]
                for column, j in column_indices:
                    numeric_close(values[column], expected[i, j])
                    matrix_cells += 1
                n_rows += 1
            assert n_rows == 462
    checked_animal, checked_group = set(), set()
    for sheet_name in ("Animal_Means", "Group_Means"):
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(iterator)
        for field in (*SOURCE_FIELDS, "OfficialFullName", "TargetStatus", "TargetHemisphere"):
            assert field in headers
        for values in iterator:
            assert len(values) <= len(headers)
            values = values + (None,) * (len(headers) - len(values))
            row = dict(zip(headers, values))
            j = target_positions[row["TargetID"]]
            source = row["Subregion"]
            if sheet_name == "Animal_Means":
                key = row["AnimalID"], source
                selected, computable, count, frequency, length = expected_animal[key]
                assert (key, j) not in checked_animal
                checked_animal.add((key, j))
            else:
                selected, computable, count, frequency, length, n_endpoint_animals, n_axon_animals = expected_group[source]
                assert integer(row["NEndpointAnimals"]) == n_endpoint_animals and integer(row["NAxonAnimals"]) == n_axon_animals
                assert (source, j) not in checked_group
                checked_group.add((source, j))
            assert integer(row["NSelected"]) == selected and integer(row["NEndpointEligible"]) == computable
            numeric_close(row["CandidateEndpointMean"], count[j])
            numeric_close(row["EndpointNeuronFrequency"], frequency[j])
            numeric_close(row["AxonTemplateLengthMeanMm"], length[j])
            assert row["OfficialFullName"] == targets[j]["OfficialFullName"]
    assert len(checked_animal) == len(animal_groups) * len(targets)
    assert len(checked_group) == len({row["Subregion"] for row in summary}) * len(targets)
    workbook.close()
    report = {"status": "passed", "n_selected": 462, "n_endpoint_eligible": 429,
              "target_level_status_columns": len(targets), "sparse_rows": len(sparse),
              "actual_atlas_target_definitions_and_volumes_checked": len(observed_targets),
              "matrix_values_checked": matrix_cells, "animal_target_rows_checked": len(checked_animal),
              "equal_animal_target_rows_checked": len(checked_group),
              "export_provenance_sha256": sha(provenance_path), "workbook_sha256": sha(workbook_path),
              "review_script_sha256": sha(__file__), "exporter_recorded_map_reconciliations": provenance["regional_reconciliation_maps"],
              "scope": "All saved matrices, per-neuron endpoint identities/totals, length totals, denominators and equal-animal means; NIfTI reconciliation is separately producer-recorded",
              "source_data_modified": False, "anatomical_acceptance": False}
    with RECEIPT.open("x", encoding="utf-8") as stream:
        json.dump(report, stream, indent=2)
        stream.write("\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
