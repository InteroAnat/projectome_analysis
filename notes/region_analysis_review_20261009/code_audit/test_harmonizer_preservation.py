"""Synthetic end-to-end workbook preservation regression; immutable inputs."""
from pathlib import Path
import contextlib
import hashlib
import importlib.util
import io
import os
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[3]
SCRIPT=Path(os.environ.get('PROJECTOME_HARMONIZER_SCRIPT',
    str(ROOT/'group_analysis/scripts/06_harmonize_atlas_to_manual.py')))
spec=importlib.util.spec_from_file_location('audit_harmonizer',SCRIPT)
harmonizer=importlib.util.module_from_spec(spec)
spec.loader.exec_module(harmonizer)


class HarmonizerPreservationTests(unittest.TestCase):
    def make_source(self,path):
        wb=Workbook();ws=wb.active;ws.title='Summary'
        ws.append(['SampleID','NeuronID','NeuronUID','Soma_Region_Auto',
            'Soma_Region_Refined','Soma_Region_Source','Unrelated_Formula'])
        ws.append([252385,'neuron001','252385::neuron001','CL_Ial','Ial',
            'auto_atlas_insula','=2+3'])
        ws['G2'].font=Font(bold=True,color='FF112233')
        ws['G2'].number_format='0.00'
        for measure in ('Length','Strength'):
            for level in ('','_L3'):
                for side in ('ipsi','contra'):
                    out=wb.create_sheet(f'Projection_{measure}{level}_{side}')
                    out.append(['SampleID','NeuronID','NeuronUID','Neuron_Type','Ial'])
                    out.append([252385,'neuron001','252385::neuron001','ITi',1.25])
        notes=wb.create_sheet('Notes');notes.append(['Notes']);notes['A2']='=1+1'
        notes['A2'].font=Font(bold=True,color='FF123456')
        notes['A2'].fill=PatternFill('solid',fgColor='FF00FF00')
        notes['A2'].number_format='0.000'
        notes['A2'].hyperlink='https://example.org/source'
        notes.column_dimensions['A'].width=41
        notes.sheet_state='hidden'
        validation=DataValidation(type='whole',operator='between',formula1=0,formula2=10)
        notes.add_data_validation(validation);validation.add('A3')
        wb.create_sheet('Mapping_Rule').append(['obsolete'])
        wb.create_sheet('Provenance').append(['obsolete'])
        wb.save(path);wb.close()

    def test_preserves_source_formulas_styles_notes_and_exact_membership(self):
        with tempfile.TemporaryDirectory() as d:
            source=Path(d)/'source.xlsx';destination=Path(d)/'harmonized.xlsx'
            self.make_source(source)
            before=source.read_bytes();before_hash=hashlib.sha256(before).hexdigest()
            with contextlib.redirect_stdout(io.StringIO()):
                harmonizer.harmonize(source,destination)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),before_hash)
            old=load_workbook(source);new=load_workbook(destination)
            try:
                self.assertEqual(new['Notes']['A2'].value,'=1+1')
                for coordinate in ('A1','A2'):
                    self.assertEqual(new['Notes'][coordinate]._style,old['Notes'][coordinate]._style)
                self.assertEqual(new['Notes']['A2'].hyperlink.target,old['Notes']['A2'].hyperlink.target)
                self.assertEqual(new['Notes'].sheet_state,'hidden')
                self.assertEqual(new['Notes'].column_dimensions['A'].width,41)
                self.assertEqual(str(new['Notes'].data_validations),str(old['Notes'].data_validations))
                self.assertEqual(new['Summary']['G2'].value,'=2+3')
                self.assertEqual(new['Summary']['G2']._style,old['Summary']['G2']._style)
                fields={cell.value:cell.column for cell in new['Summary'][1]}
                self.assertEqual(new['Summary'].cell(2,fields['Soma_Region_Refined']).value,'IAL')
                self.assertEqual(new['Summary'].cell(2,fields['Soma_Region_Refined_PreHarmonize']).value,'Ial')
                self.assertEqual(new['Summary'].cell(2,fields['Soma_Region_Source']).value,
                    'atlas_to_manual_harmonized_251637_rule')
                for name in old.sheetnames:
                    if name.startswith('Projection_'):
                        self.assertEqual(list(new[name].values),list(old[name].values))
                self.assertEqual(new['Mapping_Rule']['A1'].value,'atlas_leaf')
            finally:
                old.close();new.close()
            built=destination.read_bytes()
            with self.assertRaises(FileExistsError):harmonizer.harmonize(source,destination)
            self.assertEqual(destination.read_bytes(),built)
            self.assertEqual(source.read_bytes(),before)
            with self.assertRaises(ValueError):harmonizer.harmonize(source,source)
            self.assertEqual(source.read_bytes(),before)


if __name__=='__main__':unittest.main(verbosity=2)
