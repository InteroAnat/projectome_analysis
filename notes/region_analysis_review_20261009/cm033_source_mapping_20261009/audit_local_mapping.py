"""Read-only external CM033 source/mapping and candidate-norm snapshot.

Writes only this audit directory. Does not stage inputs, fit, warp or start jobs.
"""
from pathlib import Path
from collections import Counter
import csv
import datetime
import hashlib
import json
import re
import subprocess

import h5py
import nibabel as nib
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
DEB = Path('C:/Users/laika_yan/Documents/ChatGPT/deb_fmri')
REBUILD = DEB / 'alignment_review/cm033_Dsource_20261009'
SOURCE = Path('D:/OPTO_fMRI_CM')
FUNC = SOURCE / 'BIDS_data/CM033_BIDS/sub-CM033/func'
PRODUCTION = Path('D:/multimodal_fmri')


def sha(path):
    value = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            value.update(block)
    return value.hexdigest()


def windows_path(value):
    value = str(value)
    match = re.match(r'/mnt/([a-z])/(.*)', value)
    return Path(f'{match[1].upper()}:/{match[2]}') if match else Path(value)


def image_info(path):
    image = nib.load(str(path))
    return {
        'path': str(path), 'shape': list(image.shape),
        'zooms': list(map(float, image.header.get_zooms())),
        'affine': image.affine.tolist(),
        'units': list(image.header.get_xyzt_units()),
        'qform_code': int(image.header['qform_code']),
        'sform_code': int(image.header['sform_code']),
        'header_sha256': hashlib.sha256(image.header.binaryblock).hexdigest(),
    }


def main():
    destination = HERE / 'mapping_and_normalization_snapshot.json'
    if destination.exists():
        raise FileExistsError(destination)
    started = datetime.datetime.now(datetime.timezone.utc).isoformat()
    bindings = {}

    def bind(path):
        path = Path(path)
        bindings[str(path)] = sha(path)

    book = SOURCE / 'Info_tables&sheets/CM033.xlsx'
    bind(book)
    workbook = openpyxl.load_workbook(book, read_only=True, data_only=False)
    conversion = []
    for row_number, row in enumerate(workbook['Sheet1'].iter_rows(values_only=True), 1):
        if row_number > 1 and row[4] == 2 and row[5] == 'func':
            conversion.append({
                'excel_row': row_number, 'scanner_id': int(row[3]),
                'reconstruction': int(row[4]), 'task': str(row[6]).lower(),
                'modality': row[12],
            })
    workbook.close()
    by_task = {
        task: sorted({row['scanner_id'] for row in conversion if row['task'] == task})
        for task in ['opto', 'mstim']
    }
    acquisition = SOURCE / 'Info_tables&sheets/opto-fMRI_sessions.xlsx'
    bind(acquisition)
    workbook = openpyxl.load_workbook(acquisition, read_only=True, data_only=False)
    acquired = {}
    for row_number, row in enumerate(workbook['cm033.zJ1'].iter_rows(values_only=True), 1):
        if row[1] == 'EPI_GE' and re.fullmatch(r'[\d, ]+', str(row[0])):
            for scan in map(int, str(row[0]).replace(' ', '').split(',')):
                acquired[scan] = {
                    'excel_row': row_number, 'matrix': str(row[3]),
                    'segments': row[6], 'recorded_TR_ms': row[7],
                    'recorded_NR': row[11], 'stimulation': row[14],
                    'comments': row[15],
                }
    workbook.close()
    csv_path = SOURCE / 'cm033.csv'
    bind(csv_path)
    with csv_path.open(encoding='utf-8-sig', newline='') as stream:
        csv_rows = list(csv.DictReader(stream))
    curated = PRODUCTION / 'code/config/scans_cm033_20160120.tsv'
    bind(curated)
    with curated.open(encoding='utf-8-sig', newline='') as stream:
        selected = [row for row in csv.DictReader(stream, delimiter='\t')
                    if row['include_primary_glm'] == 'yes']
    mappings = []
    for row in selected:
        scan, task = int(row['scan_id_original']), row['task']
        legacy = by_task[task].index(scan) + 1
        path = FUNC / f'sub-CM033_task-{task.upper()}_run-{legacy:02d}_EPI.nii'
        original_exists = path.exists()
        if not original_exists:
            path = path.with_name('r' + path.name)
        if not path.exists():
            raise FileNotFoundError(path)
        image = nib.load(str(path))
        record = image_info(path)
        fingerprint = hashlib.sha256()
        # A content fingerprint corroborates distinct series, not scanner ID alone.
        for time in [0, image.shape[3] // 2, image.shape[3] - 1]:
            pixels = np.asarray(image.dataobj[..., time])
            fingerprint.update(np.ascontiguousarray(pixels).tobytes())
        metadata = path.with_suffix('.json')
        if not metadata.exists() and not original_exists:
            metadata = FUNC / f'sub-CM033_task-{task.upper()}_run-{legacy:02d}_EPI.json'
        if metadata.exists():
            bind(metadata)
            details = json.loads(metadata.read_text(encoding='utf-8'))
        else:
            details = {}
        sheet = next(item for item in conversion
                     if item['scanner_id'] == scan and item['task'] == task)
        acquisition_row = acquired.get(scan)
        recorded_frames = float(acquisition_row['recorded_NR']) if acquisition_row else None
        record.update({
            'scanner_id': scan, 'task': task, 'curated_side': row['acq'],
            'legacy_task_run_index': legacy, 'reconstruction': 2,
            'conversion_workbook_row': sheet['excel_row'],
            'conversion_csv_scan_reco_present': any(
                int(item['ScanID']) == scan and int(item['RecoID']) == 2
                for item in csv_rows),
            'original_unprefixed_exists': original_exists,
            'source_kind': 'local_converted_series' if original_exists else 'legacy_resliced_series',
            'source_json': str(metadata) if metadata.exists() else None,
            'SequenceName': details.get('SequenceName'),
            'three_frame_fingerprint_sha256': fingerprint.hexdigest(),
            'curated_TR_s': float(row['tr']),
            'TR_matches_curated': abs(float(image.header.get_zooms()[3]) - float(row['tr'])) < 1e-6,
            'acquisition_record': acquisition_row,
            'recorded_NR_matches_actual': recorded_frames == image.shape[3] if acquisition_row else None,
            'mapping_basis': 'reconstruction-2 task order, explicit MATLAB index method, acquisition/header corroboration',
            'individual_source_identity_acceptance': 'documented mapping; content comparison status below; physical LR unconfirmed',
        })
        mappings.append(record)
    content_checks = []
    raw_folder = DEB / 'alignment_review/cm033_stc_probe_20260924/bids/sub-cm033/ses-20160120/func'
    for scan in [16, 23, 39, 47, 57]:
        record = next(row for row in mappings if row['scanner_id'] == scan)
        raw = next(path for path in raw_folder.glob(f'*run-{scan}_bold.nii.gz'))
        local = Path(record['path'])
        bind(raw)
        bind(local)
        first = np.asanyarray(nib.load(str(raw)).dataobj)
        second = np.asanyarray(nib.load(str(local)).dataobj)
        same_shape = first.shape == second.shape
        check = {'scanner_id': scan, 'preserved_scanner_source': str(raw),
                 'local_source': str(local), 'same_array_shape': same_shape}
        if same_shape:
            delta = first.astype(np.float64) - second.astype(np.float64)
            check.update(exact_all_voxel_values_equal=bool(np.array_equal(first, second)),
                         maximum_absolute_value_difference=float(np.max(np.abs(delta))),
                         first_frame_exact=bool(np.array_equal(first[..., 0], second[..., 0])),
                         interpretation='voxel content correspondence does not establish physical orientation')
        content_checks.append(check)
        del first, second
        if same_shape:
            del delta
    for name in ['construct_GLM_scaninfo_033.m', 'construct_idinfo.m',
                 'concatenate_ids.m', 'first_level_GLM_cm033.m']:
        bind(SOURCE / 'Code/AccuMRNorm_ver_Binbin/Analysis' / name)
    norm = REBUILD / 'norm_reference/derivatives/norm/sub-cm033/ses-20160120_ana'
    for path in norm.iterdir():
        if path.is_file():
            bind(path)
    input_hashes = REBUILD / 'normalization_input_hashes.json'
    recorded = json.loads(input_hashes.read_text(encoding='utf-8'))
    input_checks = []
    for path, expected in recorded.items():
        actual = windows_path(path)
        bind(actual)
        input_checks.append({'path': str(actual), 'expected_sha256': expected,
                             'actual_sha256': bindings[str(actual)],
                             'match': bindings[str(actual)] == expected})
    norm_qc = REBUILD / 'normalization_qc'
    qc = json.loads((norm_qc / 'normalization_checks.json').read_text(encoding='utf-8'))
    independent_norm = []
    for label in ['CMT', 'NMT']:
        path = norm / f'cm033_20160120_Dsrc_taskA_ana_{label}_Warped.nii.gz'
        image = nib.load(str(path))
        values = np.asanyarray(image.dataobj)
        reference = windows_path(list(recorded)[1 if label == 'CMT' else 2])
        ref = nib.load(str(reference))
        jacobian = norm_qc / f'{label}_nonlinear_jacobian.nii.gz'
        bind(jacobian)
        jvalues = np.asanyarray(nib.load(str(jacobian)).dataobj)
        independent_norm.append({
            'label': label, 'forward': str(path),
            'grid_matches_bound_reference': image.shape == ref.shape and bool(np.allclose(image.affine, ref.affine, atol=1e-5)),
            'forward_all_finite': bool(np.isfinite(values).all()),
            'nonlinear_jacobian_all_finite': bool(np.isfinite(jvalues).all()),
            'nonlinear_jacobian_min': float(jvalues.min()),
            'nonlinear_jacobian_nonpositive': int((jvalues <= 0).sum()),
            'jacobian_scope': 'nonlinear only; excludes affine',
        })
    for path in norm_qc.glob('*.png'):
        bind(path)
    for name in ['normalization_completed.json', 'normalization_input_hashes.json',
                 'normalization_qc/normalization_checks.json', 'stage_local_sources.py',
                 'run_candidate_norm.py', 'run_spatial_rebuild.py',
                 'repair_staged_epi_sampling.py', 'acquisition_records.json']:
        bind(REBUILD / name)
    for name in ['CM033_D_SOURCE_SPATIAL_REBUILD_20261009.md',
                 'CM033_ALIGNMENT_NORMALIZATION_AUDIT_20261009.md']:
        bind(DEB / 'docs' / name)
    model = PRODUCTION / 'derivatives/glm/sub-cm033/ses-20160120/glm-mstim-pos1-short-tr3-smooth2mm'
    spm = model / 'SPM.mat'
    bind(spm)
    contrast = model / 'maps/con_0001.nii'
    statistic = model / 'maps/spmT_0001.nii'
    bind(contrast)
    bind(statistic)
    with h5py.File(spm, 'r') as file:
        vy = file['SPM/xY/VY']
        names = [''.join(chr(int(x)) for x in file[ref][()].flatten())
                 for ref in vy['fname'][()].flatten()]
        matrices = [file[ref][()].T for ref in vy['mat'][()].flatten()]
    counts = Counter(int(re.search(r'_run-(\d+)_', name)[1]) for name in names)
    one_based = np.eye(4)
    one_based[:3, 3] = 1
    con = nib.load(str(contrast))
    spm_matches = all(np.allclose(matrix @ one_based, con.affine, atol=1e-5) for matrix in matrices)
    process = subprocess.run(['wsl.exe', '-e', 'ps', '-eo', 'pid,ppid,etimes,args'],
                             capture_output=True, text=True, check=True)
    process_lines = [line.strip() for line in process.stdout.splitlines()
                     if re.search(r'cm033_Dsource_20261009/(stage_local_sources|run_spatial_rebuild|qc_normalization|qc_full_series)\.py(?:\s|$)', line)]
    result = {
        'status': 'documented_local_mapping_and_fresh_candidate_norm_verified_old_GLM_bridge_unestablished',
        'snapshot_started_utc': started,
        'snapshot_finished_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'task_scanner_ID_order_derived_from_workbook': by_task,
        'selected_mapping_count': len(mappings), 'mapping_rows': mappings,
        'scanner_ID_preserved_source_content_checks': content_checks,
        'fresh_normalization_inputs': input_checks,
        'fresh_normalization_independent_checks': independent_norm,
        'producer_normalization_QC_status': qc['status'],
        'live_process_snapshot_lines': process_lines,
        'source_manifest_exists_at_finish': (REBUILD / 'source_manifest.json').exists(),
        'spatial_rebuild_completed_exists_at_finish': (REBUILD / 'spatial_rebuild_completed.json').exists(),
        'new_fitted_GLMs_found_in_rebuild': list(map(str, REBUILD.rglob('SPM.mat'))),
        'left_native_model': {
            'SPM': str(spm), 'contrast': str(contrast), 'statistic': str(statistic),
            'runs_and_frames_in_actual_SPM': dict(sorted(counts.items())),
            'all_1472_SPM_input_matrices_match_contrast_after_one_based_conversion': spm_matches,
            'contrast_geometry': image_info(contrast),
            'all_recorded_input_paths_exist': all(Path(name).exists() for name in set(names)),
            'bridge_to_new_candidate_anatomy': 'not verified; no saved fit/input lineage connects this old SPM frame to fresh task-A candidate',
            'direct_new_norm_application_supported': False,
        },
        'root_visual_observations': {
            'actually_viewed_by_root': [str(norm_qc / 'CMT_forward_checkerboard_slices.png'),
                                      str(norm_qc / 'CMT_forward_checkerboard.png')],
            'scope': 'CMT forward only; not NMT/inverse/Jacobian image acceptance',
            'observation': 'Broad brain envelope/fissures/central structures correspond; contrast differs; residual peripheral/inferior/cerebellar/brainstem mismatch visible.',
            'physical_LR_or_fine_Ial_acceptance': False,
        },
        'source_sha256': bindings, 'external_source_writes': False,
        'new_fitting_registration_or_warping': False,
        'scientific_anatomical_acceptance': False,
    }
    destination.write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    fields = ['scanner_id', 'task', 'curated_side', 'legacy_task_run_index',
              'reconstruction', 'conversion_workbook_row', 'path',
              'source_kind', 'TR_matches_curated', 'recorded_NR_matches_actual']
    with (HERE / 'selected19_source_mapping.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(mappings)
    print(json.dumps({'mappings': len(mappings), 'content_checks': content_checks,
                      'independent_norm': independent_norm,
                      'left_model_frames': dict(counts), 'SPM_contrast_grid_match': spm_matches,
                      'process_snapshot': process_lines}, indent=2))


if __name__ == '__main__':
    main()
