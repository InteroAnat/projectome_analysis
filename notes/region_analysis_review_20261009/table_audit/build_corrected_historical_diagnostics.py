"""Create isolated corrections of verified historical table defects only."""
from datetime import datetime
import math
from pathlib import Path
import openpyxl
from audit_region_tables import ROOT,AUDIT,digest,save,rel

DEST=AUDIT/'corrected_diagnostics_20261009'
SOURCE_NAMES={
 'main_scripts/neuron_tables/251637_subset_multisheets.xlsx':'source-main_scripts_subset_multisheets_desc-523_unknownpartition.xlsx',
 'neuron_tables/251637_FULL.xlsx':'source-neuron_tables_FULL_desc-523_unknownpartition.xlsx',
 'neuron_tables/251637_subset.xlsx':'source-neuron_tables_subset_desc-523_unknownpartition.xlsx',
 'neuron_tables/251637_Unknown.xlsx':'source-neuron_tables_Unknown_desc-523_unknownpartition.xlsx'}
META={'NeuronID','Neuron_Type','SampleID','NeuronUID'}


def header(ws):return [c.value for c in ws[1]]


def target_row(ws):
    h=header(ws);j=h.index('NeuronID')+1
    found=[i for i in range(2,ws.max_row+1) if ws.cell(i,j).value=='523.swc']
    if len(found)!=1:raise ValueError('523 row identity must be exact and unique')
    return found[0]


def provenance(w,source,sha,rule):
    if 'Diagnostic_Provenance' in w:raise ValueError('Unexpected existing diagnostic sheet')
    ws=w.create_sheet('Diagnostic_Provenance');ws.append(['field','value'])
    for k,v in {'source_path':rel(source),'source_sha256':sha,'created_at':datetime.now().astimezone().isoformat(),'status':'isolated_diagnostic_not_canonical_or_anatomically_promoted','rule':rule,'source_unit':'stored numerical units retained; historical Summary unit undeclared','historical_523_single_partition_sum':1164.580,'fresh_current_rule_retained_projection':1156.749,'unresolved_source_lineage_difference':7.831,'hemisphere_policy':'523 historical Unknown retained; no substitution of current geometric mask R','other_rows':'preserved without new scientific validation'}.items():ws.append([k,v])


def create_523(source,destination):
    sha=digest(source);w=openpyxl.load_workbook(source,data_only=False)
    changes=set();added=[];partitions=[]
    for suffix in ('','_L3'):
        ln=f'Projection_Length{suffix}';sn=f'Projection_Strength{suffix}'
        li,lc=w[ln+'_ipsi'],w[ln+'_contra'];hi,hc=header(li),header(lc);ri,rc=target_row(li),target_row(lc)
        ai={c:li.cell(ri,j+1).value for j,c in enumerate(hi) if c not in META}
        ac={c:lc.cell(rc,j+1).value for j,c in enumerate(hc) if c not in META}
        values={}
        for c in dict.fromkeys(list(ai)+list(ac)):
            a,b=ai.get(c,0),ac.get(c,0)
            a=0 if a is None else a;b=0 if b is None else b
            if not isinstance(a,(int,float)) or not isinstance(b,(int,float)) or abs(a-b)>1e-9:raise ValueError('Historical bilateral partition values disagree')
            if a<0 or not math.isfinite(a):raise ValueError('Invalid stored projection length')
            values[c]=a
        meta=[c for c in hi if c in META]
        meta_vals={c:li.cell(ri,hi.index(c)+1).value for c in meta}
        for stem in (ln,sn):
            name=stem+'_unknown'
            if name in w:raise ValueError('Unknown diagnostic sheet already exists')
            ss=w.create_sheet(name);added.append(name)
            ss.append(meta+list(values))
            ss.append([meta_vals[c] for c in meta]+[v if stem==ln else round(math.log10(v+1),4) for v in values.values()])
        for name in (ln+'_ipsi',ln+'_contra',sn+'_ipsi',sn+'_contra'):
            ss=w[name];h=header(ss);row=target_row(ss)
            for j,c in enumerate(h,1):
                if c not in META:
                    if ss.cell(row,j).value!=0:changes.add((name,row,j))
                    ss.cell(row,j).value=0
        partitions.append({'level':'finest' if not suffix else 'L3','source_single_partition_sum':sum(values.values()),'corrected_ipsi_sum':0,'corrected_contra_sum':0,'corrected_unknown_sum':sum(values.values()),'M1_or_ancestor_length':values['M1' if not suffix else 'M1/PM'],'M1_or_ancestor_strength':round(math.log10(values['M1' if not suffix else 'M1/PM']+1),4),'unknown_nonzero_targets':{k:v for k,v in values.items() if v}})
    provenance(w,source,sha,'523 only: move one preserved identical bilateral target vector to Unknown; zero both side vectors; log10 after final target aggregation.')
    w.save(destination);w.close()
    old=openpyxl.load_workbook(source,data_only=False);new=openpyxl.load_workbook(destination,data_only=False)
    unexpected=[]
    for s in old:
        for row in s:
            for cell in row:
                actual=new[s.title].cell(cell.row,cell.column).value
                if actual!=cell.value and (s.title,cell.row,cell.column) not in changes:unexpected.append((s.title,cell.coordinate))
    if unexpected:raise ValueError(f'Unexpected changed source cells: {unexpected[:10]}')
    for s in new:
        if len(header(s))!=len(set(header(s))):raise ValueError('Duplicate output header')
    for part,suffix in zip(partitions,('','_L3')):
        ss=new[f'Projection_Length{suffix}_unknown'];h=header(ss);v=[ss.cell(2,i+1).value for i,c in enumerate(h) if c not in META]
        if abs(sum(v)-part['source_single_partition_sum'])>1e-9 or abs(sum(v)-1164.580)>1e-8:raise ValueError('Unknown partition conservation failed')
        sn=new[f'Projection_Strength{suffix}_unknown'];hs=header(sn)
        for i,c in enumerate(h):
            if c not in META and sn.cell(2,hs.index(c)+1).value!=round(math.log10(ss.cell(2,i+1).value+1),4):raise ValueError('Unknown log transform mismatch')
    so=new['Summary'];hh=header(so);rr=target_row(so)
    if so.cell(rr,hh.index('Soma_Side')+1).value!='Unknown':raise ValueError('Historical hemisphere was changed')
    old.close();new.close()
    if digest(source)!=sha:raise ValueError('Original source changed')
    return {'source':rel(source),'source_sha256':sha,'derivative':rel(destination),'derivative_sha256':digest(destination),'changed_original_data_cells':len(changes),'all_other_original_cell_values_identical':True,'historical_summary_unknown_preserved':True,'added_sheets':added+['Diagnostic_Provenance'],'partitions':partitions,'readback':'passed'}


def create_unique_headers(source,destination):
    sha=digest(source);w=openpyxl.load_workbook(source,data_only=False);removed={}
    for name in ('Ipsilateral_Length','Contralateral_Length'):
        ss=w[name];h=header(ss);field=name.split('_')[0].lower()+'_total_length';ix=[i+1 for i,c in enumerate(h) if c==field]
        if len(ix)!=2:raise ValueError('Expected two identical total headers')
        if any(ss.cell(i,ix[0]).value!=ss.cell(i,ix[1]).value for i in range(2,ss.max_row+1)):raise ValueError('Duplicate total values differ')
        removed[name]={'header':field,'removed_column_1based':ix[1],'kept_column_1based':ix[0],'data_rows':ss.max_row-1};ss.delete_cols(ix[1])
    provenance(w,source,sha,'Remove only second duplicated equal total-length column in each side sheet; retain all other sheets/cells.')
    w.save(destination);w.close();old=openpyxl.load_workbook(source,data_only=False);new=openpyxl.load_workbook(destination,data_only=False)
    for s in old:
        r=removed.get(s.title,{}).get('removed_column_1based')
        expected=[tuple(v for j,v in enumerate(row,1) if j!=r) for row in s.iter_rows(values_only=True)]
        observed=list(new[s.title].iter_rows(values_only=True))
        if expected!=observed:raise ValueError('Unique-header derivative differs beyond deleted columns')
    for s in new:
        if len(header(s))!=len(set(header(s))):raise ValueError('Remaining duplicate header')
    old.close();new.close()
    if digest(source)!=sha:raise ValueError('Original changed')
    copies=[ROOT/'main_scripts/neuron_tables/251637_laterality_projections.xlsx',ROOT/'neuron_tables/251637_laterality_projections.xlsx',ROOT/'R_analysis/tables/251637_laterality_projections.xlsx']
    if not all(digest(p)==sha for p in copies):raise ValueError('Original identical-copy binding changed')
    return {'source':rel(source),'source_sha256':sha,'source_identical_copies':[{'path':rel(p),'sha256':digest(p)} for p in copies],'derivative':rel(destination),'derivative_sha256':digest(destination),'removed_columns':removed,'all_other_original_cell_values_identical':True,'remaining_headers_unique':True,'readback':'passed'}


def main():
    DEST.mkdir(exist_ok=False)
    results=[create_523(ROOT/source,DEST/name) for source,name in SOURCE_NAMES.items()]
    results.append(create_unique_headers(ROOT/'main_scripts/neuron_tables/251637_laterality_projections.xlsx',DEST/'source-laterality_projections_desc-unique_total_columns.xlsx'))
    save(DEST/'correction_provenance.json',{'status':'isolated_diagnostic_software_readback_passed','generator':rel(Path(__file__)),'generator_sha256':digest(Path(__file__)),'source_files_modified':False,'canonical_promotion':False,'registration_acceptance':False,'fresh_lookup_lineage_difference_status':'unresolved; historical stored1164.580 retained independently of fresh current-rule1156.749','results':results})
    print({'workbooks':len(results),'all_readbacks_passed':True,'originals_unchanged':True})


if __name__=='__main__':main()
