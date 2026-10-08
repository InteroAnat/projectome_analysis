"""Fresh scalar reconstruction and exact namespace/cohort table readback."""
from collections import Counter, defaultdict
import csv
import importlib.util
import json
import math
from pathlib import Path
import nibabel as nib
import numpy as np
import openpyxl
import pandas as pd
from audit_region_tables import ROOT, AUDIT, digest, save, csvout, text, rel

OUT=AUDIT/'run_20261009'


def readrows(path):
    return list(csv.DictReader(path.open(encoding='utf-8-sig')))


def main():
    checker=ROOT/'notes/region_analysis_review_20261004/rerun_actual_region_tables.py'
    spec=importlib.util.spec_from_file_location('independent_scalar_checker',checker)
    m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
    source=ROOT/'main_scripts/processed_neurons/251637/523.swc'
    atlas=ROOT/'atlas/ARM_in_NMT_v2.1_sym.nii.gz'
    mask=ROOT/'atlas/NMT_v2.1_sym/NMT_v2.1_sym/supplemental_masks/NMT_v2.1_sym_LR_plane.nii.gz'
    key=ROOT/'atlas/ARM_key_all.txt'
    table=pd.read_csv(key,sep='\t');labels=dict(zip(table.Index,table.Abbreviation))
    fresh=m.independent_graph(source,np.asanyarray(nib.load(atlas).dataobj)[...,0,5],np.asanyarray(nib.load(mask).dataobj),labels)
    hierarchy=ROOT/'atlas/CHARM_key_table_v2.csv';hh=pd.read_csv(hierarchy)
    ancestors={}
    for label in ('CL_M1','CR_M1'):
        selected=hh[hh.Level_6_abbr.eq(label)]
        ancestors[label]=[{f'Level_{i}':r[f'Level_{i}_abbr'] for i in range(1,7)} for _,r in selected.iterrows()]
    projection=fresh['projection'];m1={k:v for k,v in projection.items() if k in ('CL_M1','CR_M1')};total=sum(m1.values())
    canonical=ROOT/'group_analysis/combined/multi_monkey_INS_combined_harmonized.xlsx'
    w=openpyxl.load_workbook(canonical,read_only=True,data_only=True)
    cohort=[];outlier_tables=[]
    for s in w:
        it=s.iter_rows(values_only=True);h=next(it,());data=list(it)
        if 'NeuronID' not in h:continue
        nid=h.index('NeuronID');counter=Counter(r[nid] for r in data)
        c={'sheet':s.title,'rows':len(data),'distinct_bare_neuron_IDs':len(counter),'repeated_bare_IDs':sum(v>1 for v in counter.values()),'rows_with_repeated_bare_IDs':sum(v for v in counter.values() if v>1),'metadata_columns':list(h[:4]),'exact_523_present':any(r[nid]=='523.swc' and ('SampleID' not in h or str(r[h.index('SampleID')])=='251637') for r in data)}
        if 'NeuronUID' in h:c['distinct_composite_UIDs']=len({r[h.index('NeuronUID')] for r in data})
        if s.title=='Summary':c['sample_counts']=dict(Counter(str(r[h.index('SampleID')]) for r in data))
        cohort.append(c)
    names=w.sheetnames;w.close()
    schemas=json.loads((OUT/'workbooks_schemas.json').read_text())['schemas']
    for d in schemas:
        for s in d.get('sheets',[]):
            if s['sheet']=='Outliers':outlier_tables.append({'file':d['representative_file'],'sha256':d['sha256'],'rows':s['rows'],'columns':s['columns'],'identity_grain':'outlier_record_not_unique_neuron'})
    summaries=readrows(OUT/'summary_units_and_cohorts.csv')
    strengths=readrows(OUT/'projection_strength_checks.csv')
    canonical_checks=[r for r in strengths if r['file']==rel(canonical)]
    domain=readrows(OUT/'unstripped_domain_observations.csv')
    buckets=defaultdict(dict)
    for r in domain:buckets[(r['file'],r['sheet'])][r['column']]=r
    actual_pooling=[]
    for (file,sheet),cols in buckets.items():
        for suffix in ('PM','CM','RM','R','Pi'):
            cc=[r for c,r in cols.items() if c in ('CL_'+suffix,'CR_'+suffix) and int(r['positive_rows'])>0]
            ss=[r for c,r in cols.items() if c in ('SL_'+suffix,'SR_'+suffix) and int(r['positive_rows'])>0]
            if cc and ss:actual_pooling.append({'file':file,'sheet':sheet,'suffix':suffix,'cortical_columns':[r['column'] for r in cc],'subcortical_columns':[r['column'] for r in ss],'cortical_positive_rows_by_column':{r['column']:int(r['positive_rows']) for r in cc},'subcortical_positive_rows_by_column':{r['column']:int(r['positive_rows']) for r in ss},'interpretation':'Both distinct domains observed in this absolute sheet; stripped suffix can pool different anatomy across neurons even without same-row positive overlap.'})
    if not (OUT/'observed_cross_domain_population_pooling_risk.csv').exists():
        csvout(OUT/'observed_cross_domain_population_pooling_risk.csv',actual_pooling)
    pi=readrows(OUT/'Pi_target_observations.csv')
    save(OUT/'focused_523_and_namespace_readback.json',{
        'software_only':True,'source_bindings':[{'path':rel(p),'sha256':digest(p)} for p in (source,atlas,mask,key,hierarchy,checker,canonical)],
        'coordinate_policy':'Current historical np.rint(XYZ/250) policy rechecked only; origin and registration anatomical acceptance remain pending.',
        'source_node_readback':fresh,'523_source_identity':'251637|523.swc','hierarchy_ancestors':ancestors,
        'historical_unknown_side_diagnostic':{'source_key_lengths_M1_only':m1,'all_retained_projection_keys':projection,'all_retained_projection_length_voxel':sum(projection.values()),'bilateral_M1_sum_voxel':total,'old_M1_both_ipsi_and_contra_length':total,'old_M1_log_before_aggregation':sum(round(math.log10(v+1),4) for v in m1.values()),'expected_separate_unknown_partition_M1':{'ipsilateral_length':0,'contralateral_length':0,'unknown_laterality_length':total,'unknown_laterality_strength':round(math.log10(total+1),4)},'fine_and_L3_targets':{'finest':'M1','L3':'M1/PM'},'status':'Counterfactual correction of historical unknown-side routing; fresh mask R is separate spatial evidence, not used to relabel this historical diagnostic.'},
        'canonical353':{'file':rel(canonical),'sha256':digest(canonical),'sheets':cohort,'Outliers_sheet_present':'Outliers' in names,'strength_pairs':len(canonical_checks),'strength_mismatches':sum(int(r['mismatched_cells']) for r in canonical_checks)},
        'outlier_sheet_inventory':outlier_tables,
        'units':{'summary_workbook_content_groups':len(summaries),'with_declared_voxel_unit':sum('"voxel"' in r['Length_Unit_values'] for r in summaries),'without_unit_in_summary':sum('not_declared_in_sheet' in r['Length_Unit_values'] for r in summaries),'status':'Missing output-adjacent unit is a reporting gap; not evidence that values have wrong physical scale. Source binding is required before converting historical lengths.'},
        'Pi':{'unstripped_parainsula_column_observations':sum(r['domain']=='parainsula' for r in pi),'unstripped_pineal_column_observations':sum(r['domain']=='pineal' for r in pi),'stripped_Pi_column_observations':sum(r['domain']=='ambiguous_stripped_Pi' for r in pi),'stripped_Pi_positive_column_observations':sum(r['domain']=='ambiguous_stripped_Pi' and int(r['positive_neuron_rows'])>0 for r in pi),'observed_numeric_pineal_parainsula_merge':'not_demonstrated_no_pineal_columns_in_observed_length_exports'},
        'actual_cross_domain_observed_absolute_sheets':dict(Counter(r['suffix'] for r in actual_pooling)),
        'historical_exact_run_receipts':'Oct4 readback hashes were reverified by source_hash_readback.csv. Older four affected workbooks do not embed exact historical SWC/atlas/code hashes; fresh matching lengths explain the error but cannot establish an unrecorded original run lineage.',
        'immutable_sources_changed':False})
    print(json.dumps({'fresh523':fresh,'observed_population_domains':dict(Counter(r['suffix'] for r in actual_pooling)),'units_missing_summary':sum('not_declared_in_sheet' in r['Length_Unit_values'] for r in summaries),'outlier_sheets':len(outlier_tables)}))


if __name__=='__main__':main()
