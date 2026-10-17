"""Nine CH00DS brains, five volume views each. Same LPS camera for every animal.

Labels come from the progress workbook, not from hardcoded region lines.
NIfTI files are read only. This script does not rewrite voxels.

Version: 1.9.1 (local/unreleased; package version)
Last updated: 2026-10-17
Change history: CHANGELOG.md (this folder)
Status: candidate (display QC only; anatomical acceptance unresolved)

Run as ``python -m Visual_toolkit.montage`` from ``main_scripts``. The old
``tools/ch00ds_5view_montage.py`` path remains a compatibility wrapper.

2026-10-16 Codex (local, unreleased): select the audited 252385 shift-corrected
DS explicitly, preserve the original montage, and bind the new figure to live
source hashes and the dataset manifest. Add manifest-bound rendering of the
nine selected native-grid candidates with one compressed-volume decode each;
the original source-only CLI remains available. Rendering/contrast settings
unchanged.
"""

from __future__ import annotations

import hashlib
import json
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import nibabel as nib
import numpy as np
import openpyxl
from scipy.ndimage import gaussian_filter, map_coordinates

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SRC = Path(
    r"\\10.102.8.200\microscopy_data\fMOST\CH00DS_datasets\processed_nii"
)


def _default_source_root() -> Path:
    """Use the VisualToolkit dataset config, with a portable fallback."""
    override = os.environ.get("PROJECTOME_DS00_PROCESSED_ROOT")
    if override:
        return Path(override)
    try:
        from .dataset_config import load
        root = (load().get("defaults") or {}).get("ds00_root")
        if root:
            return Path(root) / "processed_nii"
    except Exception:
        pass
    return DEFAULT_SRC


SRC = _default_source_root()
PLOT_DIR = REPO_ROOT / "group_analysis/evolution_20261008/overview_legacy_plot_20261012"
COPY = PLOT_DIR / "ch00ds_9brain_5view_shiftcorr252385.png"
OUT = COPY  # One project QC copy; dataset folder holds intensity data only.
WORKBOOK = PLOT_DIR / "data_progress_table.xlsx"
CORRECTED_252385 = "331_252385_ch00ds_100x100x100_orient-LPS_desc-shiftcorr_v03_aligned.nii"
CORRECTED_SHA256 = "48c021689a9ddf4324f749d33400c53ebdff0381f13d304c0ae496fdec8edad1"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(8 * 1024**2), b""):
            digest.update(block)
    return digest.hexdigest()

PANELS = [
    ("936", "251637", False),
    ("605", "252383", False),
    ("331", "252385", False),
    ("900", "252718", False),
    ("945", "252527", False),
    ("797", "252790", False),
    ("948", "252384", False),
    ("631", "252714", False),
    ("747", "250432", True),
]
DEB_MONKEYS = ("936", "605", "331", "900", "945", "797", "948", "631")
# injection_plan has no header row. track.py writes plot._PLAN_BODY in this order.
PLAN_SITES = (
    "vAIC L",
    "vAIC R",
    "Cingulate",
    "IDFA L",
    "IDFA R",
    "IDFM L",
    "IDFM R",
    "IDFP L",
    "IDFP R",
)
VIEW_KEYS = ("frontal", "superior", "back", "bottom", "coronal")
VIEW_TITLES = ("Frontal", "Superior", "Posterior", "Inferior", "Coronal")
__all__ = (
    "PANELS", "DEB_MONKEYS", "VIEW_KEYS", "VIEW_TITLES",
    "sha256", "count_text", "load_workbook_rows", "row_label", "nifti_path",
    "unit", "otsu_threshold", "anatomical_axes", "view_forward_up",
    "load_windowed", "foreground_world", "composite_view", "render_animal",
    "label_fontsize", "save_montage", "labels_for_panels", "main",
)


def count_text(value) -> str:
    """Workbook count as an integer string. Empty stays an em dash."""
    if value is None:
        return "—"
    text = str(value).strip()
    if text == "" or text.lower() == "none":
        return "—"
    number = float(text)
    if number.is_integer():
        return str(int(number))
    return text


def load_workbook_rows(path: Path) -> tuple[dict[str, dict], dict[str, str]]:
    """Progress counts and DEB site lists from data_progress_table.xlsx."""
    if not path.is_file():
        raise FileNotFoundError(path)
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "progress" not in book.sheetnames or "injection_plan" not in book.sheetnames:
            raise ValueError(f"{path} is missing progress or injection_plan")
        progress = _read_progress(book["progress"])
        plans = _read_plan(book["injection_plan"])
    finally:
        book.close()
    return progress, plans


def _read_progress(sheet) -> dict[str, dict]:
    rows = sheet.iter_rows(values_only=True)
    header = ["" if cell is None else str(cell).strip() for cell in next(rows)]
    for name in ("monkey_id", "fmost_id", "selected_n", "ARM_insula_n", "insula_8_cohort"):
        if name not in header:
            raise ValueError(f"progress sheet has no {name} column")
    index = {name: header.index(name) for name in header}
    found: dict[str, dict] = {}
    for row in rows:
        if not row or row[index["monkey_id"]] is None:
            continue
        monkey = str(row[index["monkey_id"]]).strip()
        if monkey == "TOTAL":
            continue
        found[monkey] = {
            "fmost_id": str(row[index["fmost_id"]]).strip(),
            "selected_n": row[index["selected_n"]],
            "ARM_insula_n": row[index["ARM_insula_n"]],
            "insula_8_cohort": str(row[index["insula_8_cohort"]]).strip(),
        }
    return found


def _read_plan(sheet) -> dict[str, str]:
    """Sites whose injection_plan cell is filled. Empty cells are not a 1."""
    plans: dict[str, str] = {}
    for row in sheet.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        monkey = str(row[0]).strip()
        if monkey in {"Injections", "Expected cells", "Expected total", "source"}:
            continue
        if monkey == "747":
            raise ValueError("747 is on injection_plan; Other lab has no DEB plan")
        if monkey not in DEB_MONKEYS:
            raise ValueError(f"unexpected injection_plan row {monkey}")
        cells = list(row[1 : 1 + len(PLAN_SITES)])
        if len(cells) != len(PLAN_SITES):
            raise ValueError(f"{monkey} injection_plan row is shorter than the site list")
        sites = []
        for name, value in zip(PLAN_SITES, cells):
            text = "" if value is None else str(value).strip()
            if text == "":
                continue
            sites.append(name if text == "1" else f"{name} {text}")
        if not sites:
            raise ValueError(f"{monkey} has no filled injection_plan cells")
        plans[monkey] = ", ".join(sites)
    missing = [monkey for monkey in DEB_MONKEYS if monkey not in plans]
    if missing:
        raise ValueError(f"injection_plan is missing {missing}")
    # Column 3 is Cingulate. 936's printed cell is the alignment check.
    if "Cingulate 2(L+R)" not in plans["936"]:
        raise ValueError(
            "936 injection_plan does not put 2(L+R) on Cingulate; site columns are shifted"
        )
    return plans


def row_label(monkey: str, sample: str, other: bool, progress: dict, plans: dict[str, str]) -> str:
    """total = selected_n. atlas = ARM_insula_n. DEB plan only for the eight."""
    record = progress[monkey]
    if record["fmost_id"] != sample:
        raise ValueError(f"{monkey} fmost_id is {record['fmost_id']}, panel says {sample}")
    total = count_text(record["selected_n"])
    atlas = count_text(record["ARM_insula_n"])
    if atlas == "—":
        raise ValueError(f"{monkey} ARM_insula_n is empty")
    who = f"{monkey}  Other lab" if other else monkey
    lines = [who, sample, f"total {total}", f"atlas {atlas}"]
    if other:
        if record["insula_8_cohort"] != "Other lab":
            raise ValueError(f"{monkey} cohort is {record['insula_8_cohort']}, expected Other lab")
        if monkey in plans:
            raise ValueError(f"{monkey} must not carry a DEB plan")
        return "\n".join(lines)
    if monkey not in DEB_MONKEYS:
        raise ValueError(f"{monkey} is not a DEB monkey")
    if record["insula_8_cohort"] != "DEB lab":
        raise ValueError(f"{monkey} cohort is {record['insula_8_cohort']}, expected DEB lab")
    lines.append("DEB " + plans[monkey])
    return "\n".join(lines)


def nifti_path(monkey: str, sample: str, original_sources: bool = False) -> Path:
    """100x100x90, 100x100x89.56, or 100x100x100 LPS name. Read only."""
    if sample == "252385" and not original_sources:
        if monkey != "331":
            raise ValueError("252385 must belong to animal 331")
        corrected = SRC.parent / "archive/intermediate_20261016/aligned_nii" / CORRECTED_252385
        if not corrected.is_file():
            raise FileNotFoundError(corrected)
        return corrected
    matches = sorted(SRC.glob(f"{monkey}_{sample}_ch00ds_*_orient-LPS.nii"))
    if len(matches) != 1:
        raise FileNotFoundError(f"{monkey}_{sample}: expected one NIfTI, found {matches}")
    return matches[0]


def unit(v: np.ndarray) -> np.ndarray:
    n = float(np.linalg.norm(v))
    if n == 0.0:
        raise ValueError("zero vector")
    return v / n


def otsu_threshold(values: np.ndarray) -> float:
    hi = np.percentile(values, 99.5)
    x = values[values < hi]
    if x.size < 100:
        return float(np.percentile(values, 70))
    hist, edges = np.histogram(x, bins=256)
    prob = hist.astype(np.float64)
    total = prob.sum()
    if total <= 0:
        return float(np.percentile(values, 70))
    prob /= total
    omega = np.cumsum(prob)
    centers = (edges[:-1] + edges[1:]) * 0.5
    mu = np.cumsum(prob * centers)
    mu_t = mu[-1]
    denom = omega * (1.0 - omega)
    sigma = np.zeros_like(mu)
    ok = denom > 1e-12
    sigma[ok] = (mu_t * omega[ok] - mu[ok]) ** 2 / denom[ok]
    return float(centers[int(np.argmax(sigma))])


def anatomical_axes(affine: np.ndarray) -> dict[str, np.ndarray]:
    """Directions from the NIfTI affine. These files are nibabel LPS."""
    codes = nib.aff2axcodes(affine)
    if codes != ("L", "P", "S"):
        raise RuntimeError(f"expected LPS axcodes, got {codes}")
    col = affine[:3, :3]
    left = unit(col[:, 0])
    posterior = unit(col[:, 1])
    superior = unit(col[:, 2])
    return {
        "left": left,
        "right": -left,
        "posterior": posterior,
        "anterior": -posterior,
        "superior": superior,
        "inferior": -superior,
    }


def view_forward_up(axes: dict[str, np.ndarray], name: str) -> tuple[np.ndarray, np.ndarray]:
    """forward points from the camera into the volume. up is image-up."""
    if name == "frontal":
        # Face-on. Superior up.
        forward, up = axes["posterior"], axes["superior"]
    elif name == "superior":
        # Looking down. Anterior toward the top of the panel.
        forward, up = axes["inferior"], axes["anterior"]
    elif name == "back":
        # Looking at the occiput. Superior up.
        forward, up = axes["anterior"], axes["superior"]
    elif name == "bottom":
        # Looking up from below. Anterior toward the top of the panel.
        forward, up = axes["superior"], axes["anterior"]
    elif name == "coronal":
        # Three-quarter coronal volume, not a straight duplicate of frontal.
        forward = unit(0.72 * axes["posterior"] + 1.05 * axes["left"])
        up = axes["superior"]
    else:
        raise KeyError(name)
    forward = unit(forward)
    up = up - np.dot(up, forward) * forward
    up = unit(up)
    return forward, up


def load_windowed(path: Path) -> tuple[np.ndarray, np.ndarray, float, float]:
    img = nib.load(str(path))
    affine = np.asarray(img.affine, dtype=np.float64)
    anatomical_axes(affine)
    stride = max(1, int(np.ceil(max(img.shape) / 210)))
    # Decode compressed candidate inputs once, then decimate in memory.
    full = np.asarray(img.dataobj, dtype=np.float32)
    data = full[::stride, ::stride, ::stride].copy()
    del full
    aff = affine.copy()
    aff[:3, :3] *= stride
    thr = otsu_threshold(data.ravel()[::8])
    hi = float(np.percentile(data, 99.3))
    if hi <= thr + 1:
        hi = thr + 1.0
    vol = np.clip((data - thr) / (hi - thr), 0.0, 1.0)
    # Slightly stronger blur along superior (axis 2) to soften axial section lines.
    vol = gaussian_filter(vol, sigma=(0.45, 0.45, 0.95))
    print(
        f"{path.name} shape={img.shape} stride={stride} down={vol.shape} "
        f"thr={thr:.1f} hi={hi:.1f}",
        flush=True,
    )
    return vol, aff, thr, hi


def foreground_world(vol: np.ndarray, aff: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    mask = vol > 0.08
    sel = np.flatnonzero(mask)
    if sel.size < 1000:
        raise RuntimeError("foreground too small")
    if sel.size > 300_000:
        step = int(np.ceil(sel.size / 300_000))
        sel = sel[::step]
    ii, jj, kk = np.unravel_index(sel, vol.shape)
    hom = np.stack(
        [ii.astype(np.float64), jj.astype(np.float64), kk.astype(np.float64), np.ones(ii.shape[0])],
        axis=0,
    )
    world = (aff @ hom)[:3]
    w = vol[ii, jj, kk].astype(np.float64)
    centroid = (world * w[None, :]).sum(axis=1) / w.sum()
    return world, centroid


def composite_view(
    vol: np.ndarray,
    aff: np.ndarray,
    world: np.ndarray,
    centroid: np.ndarray,
    forward: np.ndarray,
    up: np.ndarray,
    projection_bounds: tuple[float, float, float, float, float, float] | None = None,
) -> np.ndarray:
    right = unit(np.cross(up, forward))
    up = unit(np.cross(forward, right))
    delta = world - centroid[:, None]
    u = right @ delta
    v = up @ delta
    depth = forward @ delta
    if projection_bounds is None:
        u0, u1 = np.percentile(u, [0.6, 99.4])
        v0, v1 = np.percentile(v, [0.6, 99.4])
        d0, d1 = np.percentile(depth, [0.3, 99.7])
        pad_u = 0.07 * (u1 - u0)
        pad_v = 0.07 * (v1 - v0)
        u0, u1 = u0 - pad_u, u1 + pad_u
        v0, v1 = v0 - pad_v, v1 + pad_v
    else:
        u0, u1, v0, v1, d0, d1 = projection_bounds
    if min(u1 - u0, v1 - v0, d1 - d0) < 5:
        raise RuntimeError("degenerate view extent")

    bin_mm = 0.36
    nu = int(np.clip(np.round((u1 - u0) / bin_mm), 150, 260))
    nv = int(np.clip(np.round((v1 - v0) / bin_mm), 150, 280))
    nd = int(np.clip(np.round((d1 - d0) / 0.55), 56, 110))
    us = np.linspace(u0, u1, nu)
    vs = np.linspace(v0, v1, nv)
    ds = np.linspace(d0, d1, nd)
    uu, vv = np.meshgrid(us, vs, indexing="xy")
    inv_aff = np.linalg.inv(aff)

    grid = np.empty((nd, nv, nu), dtype=np.float32)
    ones = np.ones(uu.shape, dtype=np.float64)
    for zi, dist in enumerate(ds):
        world_slab = (
            centroid[0] + uu * right[0] + vv * up[0] + dist * forward[0],
            centroid[1] + uu * right[1] + vv * up[1] + dist * forward[1],
            centroid[2] + uu * right[2] + vv * up[2] + dist * forward[2],
        )
        hom = np.stack([world_slab[0], world_slab[1], world_slab[2], ones], axis=-1)
        ijk = hom @ inv_aff.T
        coords = np.stack([ijk[..., 0], ijk[..., 1], ijk[..., 2]], axis=0)
        grid[zi] = map_coordinates(vol, coords.reshape(3, -1), order=1, cval=0.0).reshape(nv, nu)

    # Light from the camera, slightly above. Front of the ray is brighter.
    shade = np.linspace(1.0, 0.42, nd, dtype=np.float32)[:, None, None]
    acc = np.zeros((nv, nu), dtype=np.float32)
    trans = np.ones((nv, nu), dtype=np.float32)
    for zi in range(nd):
        dens = grid[zi]
        alpha = 1.0 - np.exp(-dens * 8.0)
        acc += trans * alpha * dens * shade[zi]
        trans *= 1.0 - np.clip(alpha, 0.0, 0.88)
    acc = np.nan_to_num(acc, nan=0.0, posinf=0.0, neginf=0.0)
    positive = acc[acc > 0]
    if positive.size < 200:
        raise RuntimeError("blank view")
    peak = float(np.percentile(positive, 99.5))
    if peak <= 0:
        raise RuntimeError("zero peak")
    acc = np.clip(acc / peak, 0.0, 1.0)
    acc = np.power(acc, 0.88)
    acc[acc < 0.012] = 0.0
    # v increases with the up vector; flip so that direction is the top row.
    acc = acc[::-1, :]
    filled = float((acc > 0.02).mean())
    print(
        f"  extent u={u1 - u0:.1f} v={v1 - v0:.1f} d={d1 - d0:.1f} mm  "
        f"bins={nu}x{nv}x{nd} fill={filled:.3f}",
        flush=True,
    )
    if filled < 0.04:
        raise RuntimeError("view nearly blank")
    return np.stack([acc, acc * 0.985, acc * 0.94], axis=-1)


def render_animal(path: Path) -> list[np.ndarray]:
    vol, aff, _thr, _hi = load_windowed(path)
    axes = anatomical_axes(aff)
    world, centroid = foreground_world(vol, aff)
    frames = []
    for name in VIEW_KEYS:
        forward, up = view_forward_up(axes, name)
        print(f"  {name}", flush=True)
        frames.append(composite_view(vol, aff, world, centroid, forward, up))
    return frames


def label_fontsize(text: str) -> int:
    lines = text.count("\n") + 1
    longest = max(len(line) for line in text.split("\n"))
    if lines >= 6 or longest > 72:
        return 8
    if lines >= 5 or longest > 48:
        return 10
    return 12


def save_montage(all_frames: list[list[np.ndarray]], labels: list[str], dest: Path,
                 title: str = "DS montage | 252385 shift-corrected v01 | no intensity correction") -> None:
    fig, axes = plt.subplots(9, 5, figsize=(18.2, 38.0), squeeze=False)
    fig.patch.set_facecolor("black")
    fig.suptitle(title,
                 color="white", fontsize=14, y=0.99)
    for r, frames in enumerate(all_frames):
        for c, frame in enumerate(frames):
            ax = axes[r, c]
            ax.set_facecolor("black")
            ax.imshow(frame, origin="upper", interpolation="bilinear", aspect="equal")
            ax.set_xticks([])
            ax.set_yticks([])
            for spine in ax.spines.values():
                spine.set_visible(False)
            if r == 0:
                ax.set_title(
                    VIEW_TITLES[c],
                    color="white",
                    fontsize=20,
                    fontname="DejaVu Sans",
                    pad=8,
                )
        label = labels[r]
        axes[r, 0].set_ylabel(
            label,
            rotation=0,
            ha="right",
            va="center",
            color="white",
            fontsize=label_fontsize(label),
            fontname="DejaVu Sans",
            labelpad=28,
            linespacing=1.15,
        )
    fig.subplots_adjust(
        left=0.30, right=0.995, top=0.955, bottom=0.012, wspace=0.04, hspace=0.22
    )
    dest.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(dest, dpi=120, facecolor=fig.get_facecolor())
    plt.close(fig)


def labels_for_panels(cohort: dict | None = None, original_sources: bool = False) -> list[str]:
    progress, plans = load_workbook_rows(WORKBOOK)
    labels = []
    for monkey, sample, other in PANELS:
        label = row_label(monkey, sample, other, progress, plans)
        if sample == "252385" and not original_sources:
            repaired = cohort is None or any(
                x["sample"] == sample and "sourceRepair" in x["candidate_image"]
                for x in cohort["samples"])
            label += "\nDS partial source repair v03" if repaired else "\nDS shift-corrected v01"
        if cohort is not None:
            label += "\nNative-grid candidate"
            if sample == "252385":
                label += "\nshift-corrected input"
            elif sample == "250432":
                label += "\nj391 block hypothesis"
            elif sample == "252790":
                label += "\nPROVISIONAL: support caveat"
        labels.append(label)
    return labels


def main() -> None:
    global SRC, WORKBOOK
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", nargs="?", choices=["labels"])
    parser.add_argument("--source-root", type=Path,
                        help="CH00DS processed_nii directory (default: configured UNC source)")
    parser.add_argument("--workbook", type=Path,
                        help="Progress workbook used for montage row labels")
    parser.add_argument("--cohort-manifest", type=Path,
                        help="Render selected native-grid candidates in the DS cohort manifest")
    parser.add_argument("--original-sources", action="store_true",
                        help="Render all original conversions, including displaced 252385, as a hash-bound baseline")
    parser.add_argument("--output", type=Path, help="Fresh montage PNG output path")
    parser.add_argument("--receipt", type=Path, help="Fresh JSON receipt path")
    args = parser.parse_args()
    if args.original_sources and args.cohort_manifest:
        parser.error("--original-sources cannot be combined with --cohort-manifest")
    if args.source_root:
        SRC = args.source_root.resolve()
    if args.workbook:
        WORKBOOK = args.workbook.resolve()
    dest = args.output.resolve() if args.output else OUT
    receipt = args.receipt.resolve() if args.receipt else dest.with_suffix(".json")
    cohort = None
    cohort_bytes = None
    candidates = {}
    if args.cohort_manifest:
        cohort_path = args.cohort_manifest.resolve()
        cohort_bytes = cohort_path.read_bytes()
        cohort = json.loads(cohort_bytes)
        candidates = {x["sample"]: x for x in cohort["samples"]}
        expected_samples = {sample for _monkey, sample, _other in PANELS}
        if set(candidates) != expected_samples:
            raise ValueError("Cohort manifest must bind exactly the nine montage samples")
    labels = labels_for_panels(cohort, original_sources=args.original_sources)
    for (monkey, _sample, _other), text in zip(PANELS, labels):
        print("LABEL " + monkey + " " + text.replace("\n", " | "), flush=True)
    if args.command == "labels":
        return

    if dest.exists() or receipt.exists():
        raise FileExistsError("Preserve existing montage/receipt; choose a fresh output path")
    manifest_path = SRC.parent / "manifest.json"
    manifest_bytes = manifest_path.read_bytes() if cohort is None else b""
    manifest = json.loads(manifest_bytes) if manifest_bytes else {}
    expected = {item["new_relative"]: item["sha256"] for item in manifest.get("files", [])
                if item.get("new_relative") and item.get("sha256")}
    repair = manifest.get("repair_252385_v03")
    if repair:
        expected[repair["source_relative"]] = repair["source_sha256"]
    inputs = []
    all_frames = []
    for monkey, sample, _other in PANELS:
        candidate = candidates.get(sample)
        if candidate:
            cohort_root = args.cohort_manifest.resolve().parent
            path = (cohort_root / candidate["candidate_image"]).resolve()
            source_path = candidate["source_path"]
            source_hash = candidate["source_sha256"]
            run_manifest_path = (cohort_root / candidate["run_manifest"]).resolve()
            run_manifest_bytes = run_manifest_path.read_bytes()
            run_manifest = json.loads(run_manifest_bytes)
            run_hash = hashlib.sha256(run_manifest_bytes).hexdigest()
            expected_candidate_hash = candidate["candidate_image_sha256"]
            output_key = Path(candidate["candidate_image"]).relative_to(
                Path(candidate["run_directory"])).as_posix()
            if run_hash != candidate["run_manifest_sha256"]:
                raise ValueError(f"Selected run manifest SHA differs from cohort manifest: {sample}")
            if run_manifest["outputs"][output_key]["sha256"] != expected_candidate_hash:
                raise ValueError(f"Selected candidate SHA differs from run manifest: {sample}")
            if source_hash != run_manifest["source"]["sha256"]:
                raise ValueError(f"Selected candidate source SHA differs from run manifest: {sample}")
            key = candidate["candidate_image"]
        else:
            path = nifti_path(monkey, sample, original_sources=args.original_sources)
            key = path.relative_to(SRC.parent).as_posix()
            source_path = str(path)
            source_hash = None
            run_manifest_path = None
            run_manifest = None
            run_hash = None
            expected_candidate_hash = None
        before = path.stat()
        actual_hash = sha256(path)
        if candidate:
            if actual_hash != expected_candidate_hash:
                raise ValueError(f"Selected candidate SHA differs from cohort manifest: {path}")
        else:
            if expected.get(key) != actual_hash:
                raise ValueError(f"DS source hash differs from current manifest: {path}")
            if sample == "252385" and not args.original_sources and actual_hash != CORRECTED_SHA256:
                raise ValueError("252385 is not the independently audited corrected DS")
        header = nib.load(str(path))
        all_frames.append(render_animal(path))
        if cohort and (list(header.shape) != candidate["shape"] or
                       not np.allclose(header.affine, np.asarray(run_manifest["source"]["affine"]))):
            raise ValueError(f"Candidate changed the recorded native grid: {sample}")
        after = path.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            raise ValueError(f"Input changed while rendering: {path}")
        item = {"monkey": monkey, "sample": sample, "path": str(path),
                "sha256": actual_hash, "shape": list(header.shape),
                "spacing_mm": list(map(float, header.header.get_zooms())),
                "affine": header.affine.tolist(), "source_path": source_path,
                "source_sha256": source_hash,
                "correction": (run_manifest["source"]["correction"] if candidate else
                               "v01 shift retained; j417 rebuilt from aligned native sections12513-12540, max20x20x28; corrupt12511/12512 excluded")
                if sample == "252385" and not args.original_sources else "selected source/candidate on existing converted grid"}
        if candidate:
            item.update({"run_directory": candidate["run_directory"],
                         "run_manifest": str(run_manifest_path),
                         "run_manifest_sha256": run_hash,
                         "processing_status": candidate["status"],
                         "candidate_status": candidate["output_status"],
                         "boundary_decision": candidate["boundary_decision"],
                         "notches_applied": candidate["notches_applied"],
                         "block_model": candidate["block_model"]})
        inputs.append(item)
        print(f"VERIFIED {sample} {actual_hash}", flush=True)
    title = ("Native-grid intensity candidates (not registered) | 252385 shift-corrected | 252790 provisional"
             if cohort else "DS montage | 252385 partial source repair v03 | no intensity correction")
    if args.original_sources:
        title = "Original converted sources | no geometric or intensity correction"
    save_montage(all_frames, labels, dest, title)
    record = {"run_id": "DS-montage-preprocessed-cohort-20261016" if cohort else "DS-montage-shiftcorr252385-20261016", "agent": "Codex",
              "created": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
              "inputs": inputs, "labels": labels,
              "cohort_manifest": ({"path": str(args.cohort_manifest.resolve()),
                                   "sha256": hashlib.sha256(cohort_bytes).hexdigest()}
                                  if cohort else None),
              "dataset_manifest": ({"path": str(manifest_path),
                                   "sha256": hashlib.sha256(manifest_bytes).hexdigest()}
                                  if not cohort else None),
              "workbook": {"path": str(WORKBOOK), "sha256": sha256(WORKBOOK)},
              "script": {"path": str(Path(__file__).resolve()), "sha256": sha256(Path(__file__))},
              "display": {"views": VIEW_KEYS, "cameras": "existing nominal LPS axes",
                          "source_stride": "ceil(max(shape)/210)",
                          "window": "per-volume Otsu to percentile 99.3",
                          "gaussian_sigma_display_voxels": [0.45, 0.45, 0.95],
                          "rendering": "existing ray compositing; per-view 99.5 percentile and gamma 0.88",
                          "intensity_correction": bool(cohort)},
              "output": {"path": str(dest), "sha256": sha256(dest)},
              "limits": "Display rendering only; source orientation/calibration and anatomical acceptance unresolved",
              "validation": ("all nine selected candidate SHA256s match cohort and run manifests; source hashes match the run manifests; candidate size/mtime stable while read"
                             if cohort else "all nine live DS SHA256s match dataset manifest; source size/mtime stable while read")}
    if args.original_sources:
        record["run_id"] = "DS-montage-original-sources-20261016"
        record["validation"] = "All nine live original source SHA256s match dataset manifest; source size/mtime stable while read"
    receipt.write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")
    print(f"WROTE {dest}", flush=True)


if __name__ == "__main__":
    main()
