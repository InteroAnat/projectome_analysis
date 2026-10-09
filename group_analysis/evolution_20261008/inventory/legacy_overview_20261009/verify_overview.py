"""Read back the monkey view directly from sources without importing its producer."""
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]


def main():
    output = HERE / "independent_readback.json"
    if output.exists():
        raise FileExistsError("Preserve the saved independent receipt")
    bindings = {}

    def bind(path, expected=None):
        path = Path(path).resolve(strict=True)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if expected is not None and digest != expected:
            raise ValueError(f"Changed source: {path}")
        bindings[str(path)] = digest

    def read_csv(path):
        bind(path)
        with Path(path).open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    def truth(value):
        return str(value).lower() in ("1", "true")

    def count(value):
        return None if value == "" else int(float(value))

    provenance = json.loads((HERE / "provenance.json").read_text())
    bind(HERE / "provenance.json")
    for path, expected in provenance["source_sha256"].items():
        bind(path, expected)
    for name, expected in provenance["artifacts"].items():
        bind(HERE / name, expected)
    rows = read_csv(HERE / "per_monkey_insula_inventory.csv")
    inventory = HERE.parent / "live_inventory_20261008"
    samples = {r["sample"]: r for r in read_csv(inventory / "sample_inventory.csv")}
    neurons = read_csv(inventory / "neuron_inventory.csv")
    selected = read_csv(ROOT / "notes/region_analysis_review_20261009/hierarchy_tables_20261009/combined_arm_projection_tables_462/neuron_summary.csv")
    legacy_manifest = {r["animal"]: r for r in read_csv(ROOT / "group_analysis/docs/dataset_status_manifest.csv")}
    overview_path = ROOT / "group_analysis/docs/monkey_experiment_data_summary.xlsx"
    bind(overview_path)
    workbook = load_workbook(overview_path, read_only=True, data_only=True)
    overview = {str(int(r[1])): r for r in workbook["Monkey Data"].iter_rows(min_row=5, values_only=True)
                if isinstance(r[1], (int, float))}
    workbook.close()
    combined_path = ROOT / "group_analysis/combined/multi_monkey_INS_combined.xlsx"
    bind(combined_path)
    workbook = load_workbook(combined_path, read_only=True, data_only=True)
    values = workbook["Summary"].iter_rows(values_only=True)
    headers = list(next(values))
    combined = list(values)
    sid_column = headers.index("SampleID")
    combined_counts = {}
    for row in combined:
        sid = str(int(row[sid_column]))
        combined_counts[sid] = combined_counts.get(sid, 0) + 1
    workbook.close()
    henry_path = ROOT / "R_analysis/tables/somainfo_Henry_2026.04.03.xlsx"
    bind(henry_path)
    workbook = load_workbook(henry_path, read_only=True, data_only=True)
    annotations = [r for r in workbook.worksheets[0].iter_rows(min_row=2, values_only=True)
                   if re.fullmatch(r"[LR]-(IAL|IDM|IDD5|IAPM|IDV)", str(r[2]))]
    manual_ids = {int(r[1]) for r in annotations}
    workbook.close()
    assert len(annotations) == 261 and len(manual_ids) == 260
    assert len(rows) == len({r["monkey_id"] for r in rows}) == 8
    cells = 0
    for row in rows:
        animal, sid = row["monkey_id"], row["fmost_id"]
        source = overview[animal]
        assert sid == str(source[2]) == legacy_manifest[animal]["fmost_id"]
        assert row["overview_injection_sites"] == source[3]
        assert row["manifest_injection_sites"] == legacy_manifest[animal]["injection_sites"]
        subset = [r for r in neurons if r["sample"] == sid and not truth(r["excluded"])]
        atlas = [r for r in subset if truth(r["atlas_INS"])]
        candidate = [r for r in subset if truth(r["potential_INS_candidate"])]
        selected_subset = [r for r in selected if r["sample"] == sid]
        expected = {"snapshot_listed_reconstructions_n": count(samples[sid]["n_listed_reconstructions"]),
                    "snapshot_atlas_INS_n": len(atlas),
                    "spatial_candidate_lower_bound_n": len(candidate),
                    "spatial_candidate_not_atlas_INS_n": sum(not truth(r["atlas_INS"]) for r in candidate),
                    "selected_projectome_n": len(selected_subset),
                    "legacy_reconstruction_claim_n": count(legacy_manifest[animal]["ion_n"]),
                    "legacy_review_claim_n": count(legacy_manifest[animal]["tracked_insula_corrected"]),
                    "legacy_current_combined_n": combined_counts.get(sid, 0),
                    "henry_visual_INS_in_selected_n": 260 if sid == "251637" else None}
        for field, value in expected.items():
            assert count(row[field]) == value, (sid, field)
            cells += 1
        if sid == "251637":
            selected_manual = {int(re.search(r"\d+", r["neuron_id"])[0]) for r in selected_subset
                               if truth(r["Henry_coarse_INS_visual_evidence"])}
            assert manual_ids == selected_manual
            assert row["legacy_step1_atlas_INS_n"] == ""
        assert truth(row["local_verified_copy"]) == (samples[sid]["copied5um_status"] == "verified_copied_series")
        assert row["dataset_availability_elsewhere"] == "unknown"
        assert not truth(row["anatomical_acceptance"])
    priorities = sorted(rows, key=lambda r: (truth(r["local_verified_copy"]),
                       -(count(r["henry_visual_INS_in_selected_n"]) or 0), -count(r["snapshot_atlas_INS_n"]),
                       -count(r["spatial_candidate_lower_bound_n"]), r["monkey_id"]))
    assert [count(r["review_priority_rank"]) for r in priorities] == list(range(1, 9))
    difference_rows = read_csv(HERE / "legacy_reconciliation.csv")
    assert len(difference_rows) == provenance["reconciliation_rows"] == 15
    assert sum(count(r["selected_projectome_n"]) for r in rows) == 462
    tests = subprocess.run([sys.executable, "-X", "utf8", "-B", "-m", "unittest", "discover", "-s", "tests",
                            "-p", "test_monkey_insula_inventory.py", "-v"], cwd=ROOT, capture_output=True,
                           text=True, encoding="utf-8")
    log = HERE / "focused_tests.log"
    with log.open("x", encoding="utf-8") as stream:
        stream.write(tests.stdout + tests.stderr)
    assert tests.returncode == 0 and "Ran 8 tests" in tests.stderr
    bind(log)
    bind(Path(__file__))
    bind(ROOT / "tests/test_monkey_insula_inventory.py")
    receipt = {"status": "independent_legacy_monkey_overview_readback_passed",
               "checked_utc": datetime.now(timezone.utc).isoformat(), "producer_imported": False,
               "monkeys": 8, "numeric_count_cells_checked": cells, "reconciliation_rows": 15,
               "Henry_INS_annotation_rows": 261, "Henry_unique_INS_neurons": 260,
               "Henry_duplicate_fine_label_conflict_neuron": 114,
               "selected_neurons": 462, "focused_tests": 8,
               "actual_PNG_visual_review": "Codex root inspected the produced table image: readable full headers, all eight rows and both interpretation notes; no clipping observed.",
               "scientific_anatomical_acceptance": False, "source_sha256": bindings}
    with output.open("x", encoding="utf-8") as stream:
        json.dump(receipt, stream, indent=2)
        stream.write("\n")
    print(json.dumps({"status": receipt["status"], "checked_count_cells": cells, "bindings": len(bindings)}))


if __name__ == "__main__":
    main()
