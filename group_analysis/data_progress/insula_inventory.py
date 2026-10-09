"""Reconcile the legacy monkey overview with saved projectome evidence.

This is a per-monkey view, not a new cohort. Legacy claims stay separate from
snapshot atlas counts, coarse visual annotations and spatial candidates.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import manifest, paths, sources, table


def read_experiment_overview(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Monkey Data"]
    rows = []
    for cells in sheet.iter_rows(min_row=5):
        values = [cell.value for cell in cells]
        if not isinstance(values[1], (int, float)):
            continue
        rows.append({"monkey_id": str(int(values[1])), "fmost_id": str(values[2]),
                     "overview_injection_sites": values[3],
                     "overview_analysis_target_n": values[4] if isinstance(values[4], (int, float)) else None,
                     "overview_data_status": values[5],
                     "overview_source_cells": f"Monkey Data!B{cells[1].row}:F{cells[1].row}"})
    workbook.close()
    result = pd.DataFrame(rows)
    require_unique(result, ["monkey_id"], "experiment overview")
    require_unique(result, ["fmost_id"], "experiment overview")
    return result


def require_unique(frame: pd.DataFrame, columns: list[str], name: str):
    if frame[columns].isna().any().any() or frame.duplicated(columns).any():
        raise ValueError(f"Missing or duplicate exact identity in {name}")


def true(values: pd.Series) -> pd.Series:
    return values.fillna("").astype(str).str.lower().isin(["true", "1"])


def number(value):
    return None if pd.isna(value) or str(value).strip() == "" else int(value)


def build_monkey_view(overview, legacy, samples, neurons, selected, step1_paths):
    """Join by exact IDs; never merge historical aliases or channel variants."""
    require_unique(legacy, ["monkey_id"], "legacy progress table")
    require_unique(samples, ["sample"], "sample inventory")
    require_unique(neurons, ["uid"], "neuron inventory")
    require_unique(selected, ["uid"], "selected ledger")
    if set(overview.monkey_id) != set(legacy.monkey_id):
        raise ValueError("Legacy overview and progress monkey identities disagree")
    output, differences = [], []
    for source in overview.to_dict("records"):
        animal, sid = source["monkey_id"], source["fmost_id"]
        progress = legacy[legacy.monkey_id == animal].iloc[0]
        if str(progress.fmost_id) != sid:
            raise ValueError(f"Conflicting fMOST identity for monkey {animal}")
        matches = samples[samples['sample'] == sid]
        if len(matches) != 1:
            raise ValueError(f"Unmatched exact sample {sid}")
        sample = matches.iloc[0]
        current = neurons[neurons['sample'] == sid]
        eligible = current[~true(current.excluded)]
        atlas = eligible[true(eligible.atlas_INS)]
        candidate = eligible[true(eligible.potential_INS_candidate)]
        for field, actual in [("n_atlas_INS", len(atlas)), ("n_potential_INS_candidate_lower_bound", len(candidate))]:
            if number(sample[field]) != actual:
                raise ValueError(f"Independent snapshot count disagrees for {sid}: {field}")
        ledger = selected[selected['sample'] == sid]
        if len(ledger) and not ledger.registry_animal.astype(str).eq(animal).all():
            raise ValueError(f"Selected ledger animal identity disagrees for {sid}")
        manual = ledger[true(ledger.Henry_coarse_INS_visual_evidence)]
        copied = sample.copied5um_status == "verified_copied_series"
        row = {**source, "manifest_injection_sites": progress.injection_sites,
               "legacy_reconstruction_claim_n": number(progress.ion_n_traced),
               "legacy_review_claim_n": number(progress.insula_corrected_n),
               "legacy_current_combined_n": number(progress.insula_in_combined_n),
               "legacy_combined_source_labels": progress.insula_distribution_combined,
               "legacy_step1_atlas_INS_n": number(progress.step1_auto_insula_n) if step1_paths.get(sid) else None,
               "legacy_step1_status": "workbook_available" if step1_paths.get(sid) else "unavailable_workbook_not_zero",
               "snapshot_listed_reconstructions_n": number(sample.n_listed_reconstructions),
               "snapshot_atlas_INS_n": len(atlas),
               "henry_visual_INS_in_selected_n": len(manual) if len(manual) else None,
               "manual_evidence_status": "Henry_coarse_visual_INS" if len(manual) else "no_Henry_annotation_in_checked_sources",
               "spatial_candidate_lower_bound_n": len(candidate),
               "spatial_candidate_not_atlas_INS_n": len(candidate[~true(candidate.atlas_INS)]),
               "spatial_candidate_count_status": sample.potential_count_status,
               "selected_projectome_n": len(ledger),
               "selected_coarse_reviewed_INS_n": int(ledger.coarse_review_state.eq("reviewed_INS").sum()),
               "legacy_five_um_claim": progress.five_um_local,
               "local_overview_dataset_status": sample.copied5um_status,
               "dataset_availability_elsewhere": "unknown",
               "dataset_evidence_scope": sample.copied5um_evidence_scope,
               "local_copy_evidence_date": sample.copied5um_slice_evidence_date,
               "local_verified_copy": copied,
               "priority_reason": "already_verified_local_copy" if copied else "no_verified_local_copy",
               "anatomical_acceptance": False}
        comparisons = [
            ("listed_reconstructions", row["legacy_reconstruction_claim_n"], row["snapshot_listed_reconstructions_n"]),
            ("legacy_review_vs_current_combined", row["legacy_review_claim_n"], row["legacy_current_combined_n"]),
            ("step1_vs_snapshot_atlas_INS", row["legacy_step1_atlas_INS_n"], row["snapshot_atlas_INS_n"]),
            ("injection_claim_text", row["overview_injection_sites"], row["manifest_injection_sites"]),
            ("local_copy_claim", str(row["legacy_five_um_claim"]).lower().startswith("yes"), copied)]
        flags = []
        for field, previous, current_value in comparisons:
            if previous is None or previous != current_value:
                flags.append(field)
                differences.append({"monkey_id": animal, "fmost_id": sid, "field": field,
                                    "legacy_value": previous, "comparison_value": current_value,
                                    "status": "missing_or_source_scope_date_difference_not_auto_corrected"})
        row["reconciliation_flags"] = ";".join(flags)
        output.append(row)
    frame = pd.DataFrame(output)
    priority = sorted(output, key=lambda r: (r["local_verified_copy"],
                      -(r["henry_visual_INS_in_selected_n"] or 0), -r["snapshot_atlas_INS_n"],
                      -r["spatial_candidate_lower_bound_n"], r["monkey_id"]))
    ranks = {r['monkey_id']: i for i, r in enumerate(priority, 1)}
    frame["review_priority_rank"] = frame.monkey_id.map(ranks)
    return frame, pd.DataFrame(differences)


def render_overview(frame: pd.DataFrame, output: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    columns = [("monkey_id", "Monkey"), ("fmost_id", "fMOST"),
               ("snapshot_listed_reconstructions_n", "Listed\nreconstructions"),
               ("snapshot_atlas_INS_n", "Atlas-labelled\nINS"),
               ("henry_visual_INS_in_selected_n", "Henry visual\nINS"),
               ("spatial_candidate_lower_bound_n", "Spatial candidates\n(includes atlas INS)"),
               ("selected_projectome_n", "Selected for\nexploration"),
               ("legacy_review_claim_n", "Legacy review\nclaim"),
               ("legacy_current_combined_n", "Current legacy\ncombined rows"),
               ("local_verified_copy", "Verified local\nnominal 5 um copy")]
    values = []
    for row in frame.to_dict("records"):
        values.append([("Yes" if row[k] else "Not found here") if k == "local_verified_copy"
                       else ("Not documented" if pd.isna(row[k]) else str(int(row[k])) if isinstance(row[k], (float, int)) else str(row[k]))
                       for k, _ in columns])
    fig, ax = plt.subplots(figsize=(19, 5.8))
    ax.axis("off")
    ax.set_title("Per-monkey insula inventory — legacy overview reconciled with saved evidence", fontsize=14, loc="left", pad=24)
    display = ax.table(cellText=values, colLabels=[label for _, label in columns],
                       cellLoc="center", bbox=[0, .18, 1, .77], colWidths=[.055, .075, .095, .095, .095, .15, .1, .1, .1, .135])
    display.auto_set_font_size(False)
    display.set_fontsize(10)
    for (row, _), cell in display.get_celld().items():
        cell.set_edgecolor("#dddddd")
        cell.set_facecolor("#e3edf5" if row == 0 else "#f6f7f8" if row % 2 == 0 else "white")
        if row == 0:
            cell.set_text_props(weight="bold")
    ax.text(0, .09, "Atlas labels, Henry's coarse visual notes and spatial candidacy are separate evidence; counts overlap and must not be added.", fontsize=10, transform=ax.transAxes)
    ax.text(0, .035, "A local copy is not anatomical acceptance or isotropic 5 um resolution. Availability elsewhere is unknown. Original source claims remain intact.", fontsize=10, transform=ax.transAxes)
    fig.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError("Use a fresh output directory; preserve legacy and frozen outputs")
    inventory = paths.GROUP / "evolution_20261008/inventory/live_inventory_20261008"
    selected_path = paths.PROJECT / "notes/region_analysis_review_20261009/hierarchy_tables_20261009/combined_arm_projection_tables_462/neuron_summary.csv"
    overview_path = paths.DOCS / "monkey_experiment_data_summary.xlsx"
    now = datetime.now(timezone(timedelta(hours=8)))
    manifest_frame = manifest.load_manifest()
    overview = read_experiment_overview(overview_path)
    step1_paths = {sid: sources.find_step1_xlsx(sid) for sid in overview.fmost_id}
    source_files = [overview_path, paths.MANIFEST, paths.COMBINED_XLSX, selected_path,
                    inventory / "sample_inventory.csv", inventory / "neuron_inventory.csv",
                    inventory / "provenance.json", paths.REF_INS_XLSX]
    source_files += [p for p in step1_paths.values() if p is not None]
    source_files += list(Path(__file__).parent.glob("*.py"))
    source_files += [paths.PROJECT / "group_analysis/scripts/insula_label_set.py"]
    source_files = sorted({p.resolve() for p in source_files if p.exists()})
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    before = {str(p): sha(p) for p in source_files}
    legacy, _ = table.build_progress_table(manifest=manifest_frame, generated_at=now)
    read = lambda p: pd.read_csv(p, dtype=str).fillna("")
    frame, differences = build_monkey_view(overview, legacy, read(inventory / "sample_inventory.csv"),
                                          read(inventory / "neuron_inventory.csv"), read(selected_path), step1_paths)
    frame.insert(0, "generated_at", now.isoformat())
    if before != {str(p): sha(p) for p in source_files}:
        raise ValueError("Source changed during legacy/snapshot reconciliation")
    args.output.mkdir(parents=True)
    frame.to_csv(args.output / "per_monkey_insula_inventory.csv", index=False)
    differences.to_csv(args.output / "legacy_reconciliation.csv", index=False)
    render_overview(frame, args.output / "per_monkey_insula_overview.png")
    artifacts = {p.name: sha(p) for p in args.output.iterdir()}
    receipt = {"status": "legacy_overview_reconciled", "generated_at": now.isoformat(),
               "legacy_function": "group_analysis.data_progress.table.build_progress_table",
               "legacy_cleanup_runner_invoked": False, "network_requests": 0,
               "monkeys": len(frame), "selected_neurons": int(frame.selected_projectome_n.sum()),
               "all_catalog_sample_identities_retained_in_parent_inventory": 49,
               "samples_outside_exact_legacy_eight": sorted(set(read(inventory / "sample_inventory.csv")['sample']) - set(frame.fmost_id)),
               "reconciliation_rows": len(differences), "source_hashes_before_after_equal": True,
               "source_sha256": before, "artifacts": artifacts,
               "priority_policy": "No verified local copy first, then documented Henry count, atlas INS count, spatial lower bound; no overlapping count sum.",
               "manual_count_policy": "Henry coarse INS annotations in the selected ledger; absent evidence stays missing, never confirmed zero.",
               "spatial_candidate_policy": "Original 251637-screen lower bound includes atlas INS; additional non-atlas cases are reported separately.",
               "legacy_count_policy": "Unmodified progress function; unavailable step1 workbook is NA in the new view; legacy claims are not accepted anatomy.",
               "scientific_anatomical_acceptance": False}
    (args.output / "provenance.json").write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"monkeys": len(frame), "reconciliation_rows": len(differences), "output": str(args.output)}))


if __name__ == "__main__":
    main()
