from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .insula_inventory import read_experiment_overview
from .paths import DOCS, FIGURES, GROUP, LEGACY_FIGURE_GLOBS, LEGACY_TABLE_NAMES, PLAN_INSULA_EXPECTED, stamped_outputs
from .plot import _PLAN_BODY, plot_progress_table
from .sheet_overview import LEGACY_EIGHT
from .table import append_total_row, build_progress_table

WORKBOOK = DOCS / "monkey_experiment_data_summary.xlsx"
INSULA_SUMMARY = GROUP / "evolution_20261008" / "insula_summary_20261012" / "per_animal_insula_summary.csv"
COUNT_COLUMNS = [
    "neurons_screened",
    "ARM_insula_n",
    "spatial_only_n",
    "both_n",
    "Henry_coarse_INS_n",
]
COMPLETED_FILL = PatternFill("solid", fgColor="C6EFCE")
PENDING_FILL = PatternFill("solid", fgColor="B8CCE4")
HEADER_FILL = PatternFill("solid", fgColor="1A1A1A")
COHORT_FILL = PatternFill("solid", fgColor="D1C4E9")
JULY_PROGRESS = DOCS / "data_progress_table_20260701_1416.csv"


def _join_key(animal, sample) -> tuple[str, str]:
    """Identity key. Internal spacing is collapsed; the displayed cell is not edited."""
    return " ".join(str(animal).split()), " ".join(str(sample).split())


def _load_analysis_table() -> dict:
    """Call the evolution-pipeline neuron-table loader. Do not copy its numbers here."""
    import importlib.util

    path = GROUP / "evolution_20261008" / "insula_summary_20261012" / "analysis_table.py"
    spec = importlib.util.spec_from_file_location("insula_analysis_table", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    # Counts and cohort membership come from load_analysis_neuron_table.
    return module.load_analysis_neuron_table()


def _attach_checked_counts(progress: pd.DataFrame, analysis: dict) -> pd.DataFrame:
    """Copy recount fields and the cohort flag from the analysis-table loader."""
    index: dict[tuple[str, str], dict] = {}
    for record in analysis["rows"]:
        if str(record.get("channel_variant", "")) == "yes":
            continue
        if "250432ch2" in str(record.get("fMOST_samples", "")) and str(record.get("animal_id", "")).strip() == "747":
            raise ValueError("Saved summary assigns 250432ch2 to 747")
        key = _join_key(record.get("animal_id", ""), record.get("fMOST_samples", ""))
        if key in index:
            raise ValueError(f"Duplicate insula-summary identity {key}")
        index[key] = record
    rows = []
    for record in progress.to_dict("records"):
        if str(record.get("monkey_id")) == "TOTAL":
            rows.append(record)
            continue
        key = _join_key(record.get("monkey_id", ""), record.get("fmost_id", ""))
        hit = index.get(key)
        for column in COUNT_COLUMNS:
            record[column] = "" if hit is None else str(hit.get(column, ""))
        in_eight = bool(hit) and str(hit.get("insula_8_cohort", "")) == "DEB lab"
        record["insula_8_cohort"] = "DEB lab" if in_eight else "Other lab"
        record["cohort_bar_n"] = str(hit.get("ARM_insula_n", "")) if in_eight else ""
        screened = "" if hit is None else str(hit.get("neurons_screened", "")).strip()
        if screened and "neuron_screen_status" in record:
            record["neuron_screen_status"] = analysis["loader"]
        rows.append(record)
    return pd.DataFrame(rows)


def _apply_workbook_status(progress: pd.DataFrame, overview: pd.DataFrame) -> pd.DataFrame:
    """Monkey Data Completed/Pending wins for the eight. Sheet tracing text is used only outside them."""
    workbook = {
        _join_key(row.monkey_id, row.fmost_id): str(row.overview_data_status or "").strip()
        for row in overview.itertuples(index=False)
    }
    out = progress.copy()
    out["manifest_data_status"] = out["data_status"].astype(str)
    shown = []
    for record in out.to_dict("records"):
        key = _join_key(record.get("monkey_id", ""), record.get("fmost_id", ""))
        if key in workbook and workbook[key]:
            shown.append(workbook[key])
            continue
        tracing = str(record.get("sheet_tracing_status") or "").strip()
        shown.append(tracing)
    out["data_status"] = shown
    return out


def _count_cell(value) -> int:
    text = str(value).strip()
    if text == "" or text.lower() in {"nan", "none", "—", "-"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def _has_insula_neuron(row: dict) -> bool:
    """ARM insula or Henry INS. Spatial-only with both at zero is not an insula neuron."""
    return _count_cell(row.get("ARM_insula_n")) > 0 or _count_cell(row.get("Henry_coarse_INS_n")) > 0


def _keep_insula_animals(progress: pd.DataFrame) -> tuple[pd.DataFrame, list[dict]]:
    kept = []
    dropped = []
    for record in progress.to_dict("records"):
        if _has_insula_neuron(record):
            kept.append(record)
            continue
        dropped.append({
            "monkey_id": str(record.get("monkey_id", "")),
            "fmost_id": str(record.get("fmost_id", "")),
            "ARM_insula_n": str(record.get("ARM_insula_n", "")),
            "Henry_coarse_INS_n": str(record.get("Henry_coarse_INS_n", "")),
            "spatial_only_n": str(record.get("spatial_only_n", "")),
        })
    if not kept:
        raise ValueError("Insula filter removed every animal")
    return pd.DataFrame(kept), dropped


def _require_eight(frame: pd.DataFrame, label: str) -> None:
    got = set(zip(frame["monkey_id"].astype(str), frame["fmost_id"].astype(str)))
    missing = [pair for pair in LEGACY_EIGHT.items() if pair not in got]
    if missing:
        raise ValueError(f"{label} is missing legacy pairs {missing}")


def _require_saved_747(frame: pd.DataFrame) -> dict:
    hit = frame[(frame["monkey_id"].astype(str) == "747") & (frame["fmost_id"].astype(str) == "250432")]
    if len(hit) != 1:
        raise ValueError(f"Expected one 747/250432 row, found {len(hit)}")
    row = hit.iloc[0]
    expected = {
        "neurons_screened": "234",
        "ARM_insula_n": "8",
        "spatial_only_n": "21",
        "both_n": "2",
    }
    for column, value in expected.items():
        if str(row[column]).strip() != value:
            raise ValueError(f"747 {column} is {row[column]!r}, saved summary says {value}")
    if str(row["Henry_coarse_INS_n"]).strip():
        raise ValueError("747 has no Henry count in the saved summary")
    if frame["fmost_id"].astype(str).str.contains("250432ch2", regex=False).any():
        raise ValueError("250432ch2 entered the progress overview")
    if frame["monkey_id"].astype(str).isin(["ION2508004", "ION2508005"]).any():
        raise ValueError("A portal alias entered the progress overview")
    return {column: str(row[column]) for column in ("monkey_id", "fmost_id", *COUNT_COLUMNS, "data_status")}


def _write_table(display: pd.DataFrame, totals: dict, folder: Path) -> tuple[Path, Path]:
    csv_path = folder / "data_progress_table.csv"
    xlsx_path = folder / "data_progress_table.xlsx"
    display.to_csv(csv_path, index=False, encoding="utf-8-sig")

    book = Workbook()
    sheet = book.active
    sheet.title = "progress"
    columns = list(display.columns)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    status_at = columns.index("data_status") + 1
    cohort_at = columns.index("insula_8_cohort") + 1 if "insula_8_cohort" in columns else None
    for record in display.to_dict("records"):
        sheet.append([record[name] for name in columns])
        if str(record["monkey_id"]) == "TOTAL":
            continue
        status = str(record["data_status"]).strip()
        fill = COMPLETED_FILL if status == "Completed" else PENDING_FILL if status == "Pending" else None
        if fill is not None:
            sheet.cell(sheet.max_row, status_at).fill = fill
        if cohort_at is not None and str(record.get("insula_8_cohort", "")) == "DEB lab":
            sheet.cell(sheet.max_row, cohort_at).fill = COHORT_FILL
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30

    plan = book.create_sheet("injection_plan")
    for row in _PLAN_BODY:
        plan.append(row)
    plan.append(["source", "attached plan image 2026-10-12; 747 is not on this sheet"])
    meta = book.create_sheet("totals")
    meta.append(["key", "value"])
    for key, value in totals.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        meta.append([key, "" if value is None else str(value)])
    book.save(xlsx_path)
    return csv_path, xlsx_path


def _totals_from_workbook(path: Path) -> tuple[pd.DataFrame, dict]:
    """Read the table that was just saved. The figure is drawn from this file."""
    display = pd.read_excel(path, sheet_name="progress", dtype=str).fillna("")
    raw = pd.read_excel(path, sheet_name="totals", dtype=str).fillna("")
    totals = dict(zip(raw["key"], raw["value"]))
    for key in ("analysis_expected", "ion_traced", "insula_corrected", "insula_expected_plan", "monkeys"):
        totals[key] = int(float(totals[key]))
    return display, totals


def write_progress_overview(
    output_dir: Path,
    include_online_sheet: bool = True,
    write_figure: bool = True,
) -> dict:
    """Parse Monkey Data and Progress Details, save the table, then plot that file.

    Insula counts and cohort membership come from
    ``analysis_table.load_analysis_neuron_table`` (462-neuron ledger plus the
    saved recount). A change to that table must be passed through this
    function so ``plot_progress_table`` redraws it. The July combined
    workbook is not the reality-bar source.
    """
    output = Path(output_dir)
    if output.exists() and not output.is_dir():
        raise FileExistsError(output)
    blocked = [
        (GROUP / "evolution_20261008" / "insula_summary_20261012").resolve(),
        (GROUP / "evolution_20261008" / "inventory" / "legacy_overview_20261009").resolve(),
        DOCS.resolve(),
    ]
    resolved = output.resolve()
    if any(resolved == path or path in resolved.parents for path in blocked):
        raise ValueError(f"Refusing to write into a protected folder: {output}")

    overview = read_experiment_overview(WORKBOOK)
    pairs = set(zip(overview["monkey_id"].astype(str), overview["fmost_id"].astype(str)))
    if pairs != set(LEGACY_EIGHT.items()):
        raise ValueError(f"Monkey Data is not the eight known pairs: {sorted(pairs)}")

    progress, totals = build_progress_table(
        generated_at=datetime.now().astimezone(),
        include_online_sheet=include_online_sheet,
    )
    _require_eight(progress, "progress table")
    progress = _apply_workbook_status(progress, overview)
    # Insula counts and cohort membership: evolution loader, not the July combined workbook.
    analysis = _load_analysis_table()
    progress = _attach_checked_counts(progress, analysis)
    animal_747 = _require_saved_747(progress) if include_online_sheet else {}
    if not include_online_sheet and (progress["monkey_id"].astype(str) == "747").any():
        raise ValueError("Offline progress overview invented animal 747")
    n_eight = int(progress["insula_8_cohort"].astype(str).eq("DEB lab").sum())
    if n_eight != 8:
        raise ValueError(f"Analysis table marked {n_eight} insula-eight rows, expected 8")
    progress, dropped_no_insula = _keep_insula_animals(progress)
    if int(progress["insula_8_cohort"].astype(str).eq("DEB lab").sum()) != 8:
        raise ValueError("Insula filter dropped a DEB lab animal")
    if include_online_sheet and not (
        (progress["monkey_id"].astype(str) == "747") & (progress["fmost_id"].astype(str) == "250432")
    ).any():
        raise ValueError("Insula filter dropped 747/250432")

    july = pd.read_csv(JULY_PROGRESS, dtype=str, keep_default_na=False)
    july_total = july.loc[july["monkey_id"] == "TOTAL"].iloc[0]
    old_ion = int(float(july_total["ion_n_traced"]))
    old_insula = int(float(july_total["insula_corrected_n"]))
    plan_totals = dict(totals)
    if "cohort_monkeys" in plan_totals:
        plan_totals["monkeys"] = plan_totals["cohort_monkeys"]
    plan_totals["insula_expected_plan"] = int(plan_totals.get("insula_expected_plan", PLAN_INSULA_EXPECTED))
    plan_totals["analysis_expected_manifest_sum"] = int(plan_totals["analysis_expected"])
    # July figure expectation the user identified. Not the manifest sum 1225, and not the 2400 cell target.
    plan_totals["analysis_expected"] = 2429
    plan_totals["analysis_expected_label"] = "July plan expectation"
    # Reality bars: eight-only recount. Leave the July combined 353 out.
    plan_totals["july_ion_traced"] = old_ion
    plan_totals["july_insula_corrected"] = old_insula
    plan_totals["ion_traced"] = int(analysis["eight_screened"])
    plan_totals["insula_corrected"] = int(analysis["eight_arm_insula"])
    plan_totals["reality_ion_meaning"] = "neurons_screened for the insula eight"
    plan_totals["reality_insula_meaning"] = "ARM_insula_n for the insula eight"
    plan_totals["neuron_table_loader"] = analysis["loader"]
    plan_totals["neuron_table_ledger"] = analysis["ledger_path"]
    plan_totals["neuron_table_summary"] = analysis["summary_path"]
    status_changes = []
    for record in progress.to_dict("records"):
        if str(record.get("insula_8_cohort", "")) != "DEB lab":
            continue
        previous = str(record.get("manifest_data_status", "")).strip()
        shown = str(record.get("data_status", "")).strip()
        if previous != shown:
            status_changes.append({
                "animal": str(record["monkey_id"]),
                "fmost": str(record["fmost_id"]),
                "july_manifest_status": previous,
                "shown_status": shown,
                "source": "Monkey Data workbook",
            })
    display = append_total_row(progress, plan_totals)
    extras = progress[progress.get("in_legacy_eight", pd.Series("yes", index=progress.index)).astype(str) == "no"]
    if len(extras):
        total_at = display.index[display["monkey_id"] == "TOTAL"][0]
        display.loc[total_at, "data_status"] = f"8 plan + {len(extras)} outside"
        display.loc[total_at, "notes"] = (
            str(display.loc[total_at, "notes"])
            + f"; {len(extras)} animals outside the DEB lab eight remain after dropping rows with no insula neuron"
        )

    sheet_meta = totals.get("online_sheet") or {"fetch": "offline", "fetch_error": "", "snapshot": ""}
    source_label = sheet_meta.get("fetch") or "offline"
    if sheet_meta.get("snapshot"):
        source_label = f"{source_label}:{Path(sheet_meta['snapshot']).name}"
    display["progress_details_source"] = source_label
    plan_totals["online_fetch"] = sheet_meta.get("fetch", "")
    plan_totals["online_snapshot"] = sheet_meta.get("snapshot", "")
    plan_totals["workbook"] = str(WORKBOOK)
    plan_totals["insula_summary"] = analysis["summary_path"]
    plan_totals["progress_details_source"] = source_label

    output.mkdir(parents=True, exist_ok=True)
    for name in ("data_progress_table.csv", "data_progress_table.xlsx", "data_progress_table.png", "provenance.json"):
        (output / name).unlink(missing_ok=True)
    csv_path, xlsx_path = _write_table(display, plan_totals, output)
    figure = output / "data_progress_table.png"
    if write_figure:
        saved, saved_totals = _totals_from_workbook(xlsx_path)
        plot_progress_table(saved, saved_totals, figure)

    outside = []
    if "in_legacy_eight" in progress.columns:
        outside = (
            progress.loc[progress["in_legacy_eight"].astype(str) == "no", ["monkey_id", "fmost_id"]]
            .astype(str)
            .values.tolist()
        )
    receipt = {
        "pipeline": "group_analysis.data_progress.track.write_progress_overview",
        "plot_function": "group_analysis.data_progress.plot.plot_progress_table",
        "figure_title": "Data progress — insula multi-monkey cohort",
        "include_online_sheet": include_online_sheet,
        "online_sheet": sheet_meta,
        "workbook": str(WORKBOOK),
        "neuron_table_loader": analysis["loader"],
        "neuron_table_ledger": analysis["ledger_path"],
        "neuron_table_summary": analysis["summary_path"],
        "ledger_rows": analysis["ledger_rows"],
        "reality_bars": {
            "old_ion_traced": old_ion,
            "new_ion_traced": int(analysis["eight_screened"]),
            "old_insula_corrected": old_insula,
            "new_insula_arm": int(analysis["eight_arm_insula"]),
            "expected_ion_unchanged": int(plan_totals["analysis_expected"]),
            "plan_target_unchanged": int(plan_totals["insula_expected_plan"]),
        },
        "insula_neuron_rule": "ARM_insula_n > 0 or Henry_coarse_INS_n > 0; spatial-only is not enough",
        "progress_table_animals": int(len(progress)),
        "dropped_no_insula_neuron": dropped_no_insula,
        "status_changes_from_july_manifest": status_changes,
        "insula_summary": analysis["summary_path"],
        "row_747": animal_747,
        "animals_outside_eight": [{"monkey_id": animal, "fmost_id": sample} for animal, sample in outside],
        "anatomical_acceptance": False,
        "added_to_462_ledger": False,
    }
    (output / "provenance.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {
        "table": csv_path,
        "workbook": xlsx_path,
        "figure": figure if write_figure else None,
        "totals_dict": plan_totals,
        "online_sheet": sheet_meta,
        "row_747": animal_747,
        "animals_outside_eight": receipt["animals_outside_eight"],
        "provenance": output / "provenance.json",
    }


def remove_legacy_outputs() -> list[str]:
    """Drop superseded dataset_status figures and unstamped data_progress files."""
    removed: list[str] = []
    for pattern in LEGACY_FIGURE_GLOBS:
        for path in FIGURES.glob(pattern):
            path.unlink(missing_ok=True)
            removed.append(str(path))
    for name in LEGACY_TABLE_NAMES:
        path = DOCS / name
        if path.exists():
            path.unlink()
            removed.append(str(path))
    return removed


def run(
    write_figure: bool = True,
    cleanup_legacy: bool = True,
    include_online_sheet: bool | None = None,
    output_dir: Path | None = None,
) -> dict:
    """Build the progress table and figure.

    With ``output_dir``, this calls ``write_progress_overview`` (workbook, online
    sheet, saved insula counts, then ``plot_progress_table``). That path reads
    the online sheet unless ``include_online_sheet`` is False. Without
    ``output_dir``, the historical eight-monkey writer stays offline and does
    not touch the overview folder or the July CSV.
    """
    if output_dir is not None:
        online = True if include_online_sheet is None else include_online_sheet
        return write_progress_overview(
            output_dir,
            include_online_sheet=online,
            write_figure=write_figure,
        )
    progress, totals = build_progress_table(include_online_sheet=bool(include_online_sheet))
    display = append_total_row(progress, totals)

    out_table, out_figure = stamped_outputs(totals["file_stamp"])

    out_table.parent.mkdir(parents=True, exist_ok=True)
    display.to_csv(out_table, index=False, encoding="utf-8-sig")

    if write_figure:
        out_figure.parent.mkdir(parents=True, exist_ok=True)
        plot_progress_table(display, totals, out_figure)

    removed = remove_legacy_outputs() if cleanup_legacy else []

    return {
        "table": out_table,
        "figure": out_figure if write_figure else None,
        "totals_dict": totals,
        "removed_legacy": removed,
    }


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Build the data progress overview.")
    parser.add_argument("--output", type=Path, default=None, help="Fresh folder for the progress overview.")
    parser.add_argument("--offline", action="store_true", help="Skip the online sheet and the saved TSV.")
    args = parser.parse_args(argv)
    if args.output is not None:
        result = run(
            write_figure=True,
            cleanup_legacy=False,
            include_online_sheet=not args.offline,
            output_dir=args.output,
        )
        sheet = result.get("online_sheet") or {}
        print("Progress overview written:")
        print(f"  table:  {result['table']}")
        print(f"  xlsx:   {result['workbook']}")
        print(f"  figure: {result['figure']}")
        print(f"  sheet:  {sheet.get('fetch', 'offline')}")
        if sheet.get("snapshot"):
            print(f"  snapshot: {sheet['snapshot']}")
        return
    result = run(write_figure=True)
    t = result["totals_dict"]
    print("Data progress track updated:")
    print(f"  table:  {result['table']}")
    print(f"  figure: {result['figure']}")
    if result["removed_legacy"]:
        print("  removed legacy outputs:")
        for path in result["removed_legacy"]:
            print(f"    - {path}")
    print(
        f"  summary @ {t['generated_at']}: "
        f"insula {t['insula_corrected']}/{t['insula_expected_plan']} corrected; "
        f"ION {t['ion_traced']}/{t['analysis_expected']}; "
        f"combined {t['insula_in_combined']}; "
        f"5µm {t['five_um_ready']}/{t['monkeys']}"
    )


if __name__ == "__main__":
    main()
