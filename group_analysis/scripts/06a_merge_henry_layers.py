"""
Phase 6a: Merge Henry (2026-04-03) area + layer annotations into the
harmonized combined table for sample 251637.

Source
------
Sheet: ``251637 mostly insula``
Columns: Folder, SWC, Area, Layer

Resolved in this order: ``--henry-source``, ``PROJECTOME_HENRY_LAYERS``,
``R_analysis/tables/somainfo_Henry_2026.04.03.xlsx`` if that file is present,
then the lab share ``X:/fMOST/251637/Area and Layers by Henry_2026.04.03.xlsx``.

Join
----
SWC (int) -> NeuronID = ``{SWC:03d}.swc`` on Summary rows with SampleID=251637.

Duplicate SWC rows (currently only 114.swc) are resolved by preferring the
Henry Area that contains the harmonized ``Soma_Region_Refined`` leaf
(e.g. IDD5 vs IDM).

Outputs
-------
Atomically updates
``group_analysis/combined/multi_monkey_INS_combined_harmonized.xlsx``:
  - Summary: ``Soma_Area_Henry``, ``Cortical_Layer``, ``Layer_Source``
  - Henry_Layer_Provenance: source identity + merge audit

The workbook is edited with openpyxl rather than round-tripped through pandas,
so unrelated sheets, formulas, formatting, validation, and workbook metadata
are retained.
"""
from __future__ import annotations

import argparse
import hashlib
import os
import tempfile
from datetime import date

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = r"D:\projectome_analysis"
GROUP_DIR = os.path.join(PROJECT_ROOT, "group_analysis")
HARMONIZED_XLSX = os.path.join(
    GROUP_DIR, "combined", "multi_monkey_INS_combined_harmonized.xlsx"
)
HENRY_SHEET = "251637 mostly insula"
HENRY_LAB_XLSX = r"X:\fMOST\251637\Area and Layers by Henry_2026.04.03.xlsx"
HENRY_REPO_XLSX = os.path.join(
    PROJECT_ROOT, "R_analysis", "tables", "somainfo_Henry_2026.04.03.xlsx"
)
LAYER_SOURCE = "henry_20260403"
TARGET_COLUMNS = ("Soma_Area_Henry", "Cortical_Layer", "Layer_Source")


def _neuron_id_from_swc(swc: int | float) -> str:
    return f"{int(swc):03d}.swc"


def _pick_henry_row(
    henry_rows: pd.DataFrame, soma_region_refined: str
) -> pd.Series:
    if len(henry_rows) == 1:
        return henry_rows.iloc[0]
    region = str(soma_region_refined or "").strip()
    if region:
        match = henry_rows[
            henry_rows["Area"].astype(str).str.contains(region, regex=False, na=False)
        ]
        if len(match) == 1:
            return match.iloc[0]
        if len(match) > 1:
            return match.iloc[0]
    return henry_rows.sort_values("Area").iloc[0]


def resolve_henry_xlsx(explicit: str | None = None) -> str:
    """Pick the Henry workbook without requiring one machine's drive letter."""
    if explicit:
        return explicit
    env = os.environ.get("PROJECTOME_HENRY_LAYERS", "").strip()
    if env:
        return env
    if os.path.isfile(HENRY_REPO_XLSX):
        return HENRY_REPO_XLSX
    return HENRY_LAB_XLSX


def load_henry_table(path: str | None = None) -> pd.DataFrame:
    path = resolve_henry_xlsx(path)
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    df = pd.read_excel(path, sheet_name=HENRY_SHEET)
    required = {"SWC", "Area", "Layer"}
    missing = required - set(df.columns)
    if missing:
        raise KeyError(f"Henry sheet missing columns: {sorted(missing)}")
    df = df.copy()
    df["NeuronID"] = df["SWC"].map(_neuron_id_from_swc)
    return df


def _sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def merge_henry_layers(
    harmonized_xlsx: str = HARMONIZED_XLSX,
    henry_xlsx: str | None = None,
) -> None:
    henry_xlsx = resolve_henry_xlsx(henry_xlsx)
    if not os.path.isfile(harmonized_xlsx):
        raise FileNotFoundError(harmonized_xlsx)
    if not os.path.isfile(henry_xlsx):
        raise FileNotFoundError(henry_xlsx)

    henry = load_henry_table(henry_xlsx)

    wb = load_workbook(harmonized_xlsx)
    if "Summary" not in wb.sheetnames:
        raise KeyError("Workbook missing Summary sheet")
    ws = wb["Summary"]
    headers = _header_map(ws)
    required = {"SampleID", "NeuronID", "Soma_Region_Refined"}
    missing = required - set(headers)
    if missing:
        raise KeyError(f"Summary missing columns: {sorted(missing)}")

    for name in TARGET_COLUMNS:
        if name not in headers:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=name)
            headers[name] = col

    target_rows: list[int] = []
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, headers["SampleID"]).value == 251637:
            target_rows.append(row_idx)
            # Rebuild the scoped annotations from scratch on every run.
            for name in TARGET_COLUMNS:
                ws.cell(row_idx, headers[name]).value = None

    n_merged = 0
    duplicate_resolutions: list[str] = []
    unmatched: list[str] = []
    selected_duplicates: list[str] = []

    for row_idx in target_rows:
        nid = str(ws.cell(row_idx, headers["NeuronID"]).value)
        rows = henry[henry["NeuronID"] == nid]
        if rows.empty:
            unmatched.append(nid)
            continue
        if len(rows) > 1:
            duplicate_resolutions.append(nid)
        region = ws.cell(row_idx, headers["Soma_Region_Refined"]).value
        pick = _pick_henry_row(rows, region)
        if len(rows) > 1:
            selected_duplicates.append(f"{nid}:{pick['Area']}:{pick['Layer']}")

        ws.cell(row_idx, headers["Soma_Area_Henry"], value=pick["Area"])
        layer = pick["Layer"]
        layer_value = (
            int(layer)
            if pd.notna(layer) and float(layer).is_integer()
            else layer
        )
        ws.cell(row_idx, headers["Cortical_Layer"], value=layer_value)
        ws.cell(row_idx, headers["Layer_Source"], value=LAYER_SOURCE)
        n_merged += 1

    if unmatched or n_merged != len(target_rows):
        raise RuntimeError(
            "Henry merge coverage failure: "
            f"merged={n_merged}, expected={len(target_rows)}, "
            f"unmatched={unmatched}"
        )

    print(f"[06a] Henry rows: {len(henry)}")
    print(f"[06a] merged 251637 neurons: {n_merged} / {len(target_rows)}")
    if duplicate_resolutions:
        print(
            "[06a] duplicate SWC resolved: "
            f"{sorted(set(duplicate_resolutions))}"
        )

    if "Henry_Layer_Provenance" in wb.sheetnames:
        del wb["Henry_Layer_Provenance"]
    prov = wb.create_sheet("Henry_Layer_Provenance")
    provenance = {
        "source_file": henry_xlsx,
        "source_sha256": _sha256(henry_xlsx),
        "source_mtime": pd.Timestamp(os.path.getmtime(henry_xlsx), unit="s").isoformat(),
        "source_sheet": HENRY_SHEET,
        "source_rows": len(henry),
        "merge_date": date.today().isoformat(),
        "layer_source_tag": LAYER_SOURCE,
        "sample_id": 251637,
        "target_rows": len(target_rows),
        "neurons_merged": n_merged,
        "unmatched_ids": ", ".join(unmatched),
        "duplicate_swc_resolved": ", ".join(sorted(set(duplicate_resolutions))),
        "duplicate_rows_selected": ", ".join(selected_duplicates),
        "columns_added": ", ".join(TARGET_COLUMNS),
        "layer_resolution": "whole_layer_no_subdivision",
    }
    prov.append(list(provenance))
    prov.append(list(provenance.values()))

    out_dir = os.path.dirname(os.path.abspath(harmonized_xlsx))
    fd, temp_path = tempfile.mkstemp(
        prefix=".henry_layers_", suffix=".xlsx", dir=out_dir
    )
    os.close(fd)
    try:
        wb.save(temp_path)
        check = load_workbook(temp_path, read_only=True, data_only=False)
        check_headers = _header_map(check["Summary"])
        if not set(TARGET_COLUMNS).issubset(check_headers):
            raise RuntimeError("Atomic output validation failed: layer columns missing")
        if "Henry_Layer_Provenance" not in check.sheetnames:
            raise RuntimeError("Atomic output validation failed: provenance missing")
        check.close()
        os.replace(temp_path, harmonized_xlsx)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

    print(f"[06a] wrote {harmonized_xlsx}")


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", default=HARMONIZED_XLSX)
    parser.add_argument(
        "--henry-source",
        default=None,
        help=(
            "Henry area/layer workbook. Default search: "
            "PROJECTOME_HENRY_LAYERS, then "
            "R_analysis/tables/somainfo_Henry_2026.04.03.xlsx, then "
            r"X:\fMOST\251637\Area and Layers by Henry_2026.04.03.xlsx"
        ),
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = _parse_args()
    merge_henry_layers(args.workbook, args.henry_source)
