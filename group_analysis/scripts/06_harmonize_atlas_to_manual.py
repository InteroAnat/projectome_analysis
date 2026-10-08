"""
Phase 6: Harmonize new-monkey atlas-derived insula labels to 251637's
manual sub-region scheme (improvised mapping).

Why
---
251637's neurons were manually curated into the Evrard 2014-style scheme
(IAL/IAPM/IDD5/IDM/IDV). New monkeys (252383/252384/252385) carry the
CHARM/NMT atlas vocabulary (Ial/Iai/Iam/Iapm/Iapl/Ia-Id/Ig/Pi/...). To
keep cross-monkey strata comparable, we map atlas leaves to manual labels
using the empirical atlas <-> manual crosstab observed within 251637
itself (n=260, both labels available).

Empirical 251637 crosstab (atlas leaf -> manual leaf):
    Ial        IAL    92.1%   IAPM 7.9%
    Iai        IAPM   55.0%   IAL 45.0%        (ambiguous)
    Ia/Id      IDM    62.9%   IDD5 34.3%   IDV 2.9%   (ambiguous)
    Ig         IDD5  100.0%
    PrCO       IAL   100.0%
    Unknown_0  IDM    76.9%   IDD5 23.1%

This script only edits rows whose `Soma_Region_Source == auto_atlas_insula`
(i.e., new-monkey neurons not already coord-rescued via 251637's
sub-region bbox). 251637 manual rows and coord-rescued rows are
preserved as-is.

Important caveat
----------------
This is an *improvised* harmonization, marked accordingly in
`Soma_Region_Source`:
    atlas_to_manual_harmonized_251637_rule           - high confidence
    atlas_to_manual_harmonized_251637_rule_ambiguous - needs visual QC
    atlas_Ig_to_IG_granular_user_rule_20260926       - user decision 2026-09-26

User decision (2026-09-26): atlas ``Ig`` maps to granular ``IG`` (not IDD5).
Empirical 251637 crosstab had Ig→IDD5 at 100% (24 cells), but that collapsed
granular Ig into the dysgranular stratum; treat those as granular for cohort
refresh. Atlas ``G`` (gustatory cortex) is **not** remapped to IG — leave G as is.
Only ``Soma_Region_Source == auto_atlas_insula`` rows are edited; curated
251637 and coord-rescued rows are preserved.

Outputs
-------
Writes a sibling file
`group_analysis/combined/multi_monkey_INS_combined_harmonized.xlsx`
(override with ``PROJECTOME_HARMONIZED_OUT``; input via ``PROJECTOME_COMBINED_IN``)
with:
  - Summary: adds `Soma_Region_Refined_PreHarmonize` column,
             updates `Soma_Region_Refined` and `Soma_Region_Source`.
  - Mapping_Rule (new): documents the empirical crosstab and applied rule.
  - Provenance (regenerated): updated counts.
Untouched workbook cells retain their values, formulas and formatting through
openpyxl. Existing Phase 6a enrichment still updates its documented Henry
Summary annotation fields and Henry_Layer_Provenance sheet. The derivative
XLSX archive is reserialized; byte-for-byte equality is promised only for the
unchanged source file.

The original `multi_monkey_INS_combined.xlsx` is left untouched. Once
verified, the user can swap names manually.
"""
from __future__ import annotations

import os
import importlib.util
import tempfile
from pathlib import Path
from numbers import Real

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

PROJECT_ROOT = r"D:\projectome_analysis"
GROUP_DIR = os.path.join(PROJECT_ROOT, "group_analysis")
COMBINED_XLSX = os.environ.get(
    "PROJECTOME_COMBINED_IN",
    os.path.join(GROUP_DIR, "combined", "multi_monkey_INS_combined.xlsx"),
)
OUT_XLSX = os.environ.get(
    "PROJECTOME_HARMONIZED_OUT",
    os.path.join(
        GROUP_DIR, "combined", "multi_monkey_INS_combined_harmonized.xlsx"
    ),
)

# Atlas-leaf -> manual-leaf mapping.
# Format: atlas_leaf -> (dominant_manual, confidence_label)
# confidence "user_ig" → special Soma_Region_Source tag (granular IG).
EMPIRICAL_MAPPING = {
    "Ial":       ("IAL",  "high"),       # 92.1% -> IAL
    "PrCO":      ("IAL",  "high"),       # 100.0% -> IAL (already coord-rescued in v5)
    "Ig":        ("IG",   "user_ig"),    # user rule 2026-09-26: granular IG (not IDD5)
    "Iai":       ("IAPM", "ambiguous"),  # 55/45 IAPM/IAL split
    "Ia/Id":     ("IDM",  "ambiguous"),  # 63/34/3 IDM/IDD5/IDV split
    "Unknown_0": ("IDM",  "ambiguous"),  # 77/23 IDM/IDD5 split
}
# Do NOT map atlas "G" (gustatory) → IG.

# 251637 reference crosstab (frozen from neuron_tables_new + step1 join, n=260)
# Used in the Mapping_Rule sheet for documentation.
REFERENCE_CROSSTAB = pd.DataFrame(
    [
        # (atlas_leaf, manual_leaf, count)
        ("Ia/Id",     "IDD5",  36),
        ("Ia/Id",     "IDM",   66),
        ("Ia/Id",     "IDV",    3),
        ("Iai",       "IAL",    9),
        ("Iai",       "IAPM",  11),
        ("Ial",       "IAL",   58),
        ("Ial",       "IAPM",   5),
        ("Ig",        "IDD5",  24),
        ("PrCO",      "IAL",   35),
        ("Unknown_0", "IDD5",   3),
        ("Unknown_0", "IDM",   10),
    ],
    columns=["atlas_leaf", "manual_leaf", "n"],
)


def _strip_prefix(label: str) -> str:
    """Drop the side prefix (CL_/CR_) to get the raw atlas leaf."""
    if not isinstance(label, str):
        return ""
    for pre in ("CL_", "CR_"):
        if label.startswith(pre):
            return label[len(pre):]
    return label


def _sample_token(value):
    if pd.isna(value):
        raise ValueError("SampleID must be present")
    if isinstance(value, Real) and float(value).is_integer():
        return str(int(value))
    token = str(value).strip()
    if not token:
        raise ValueError("SampleID must be present")
    return token


def _ordered_uids(frame, sheet_name):
    required = {"SampleID", "NeuronID"}
    if not required.issubset(frame.columns):
        raise ValueError(f"{sheet_name} missing identity columns: {sorted(required - set(frame.columns))}")
    if frame.NeuronID.isna().any():
        raise ValueError(f"{sheet_name} has missing NeuronID")
    pairs = []
    for sample, neuron in zip(frame.SampleID, frame.NeuronID):
        neuron = str(neuron).strip()
        if not neuron:
            raise ValueError(f"{sheet_name} has blank NeuronID")
        pairs.append((_sample_token(sample), neuron))
    if len(pairs) != len(set(pairs)):
        raise ValueError(f"{sheet_name} has duplicate full neuron identities")
    if "NeuronUID" in frame:
        expected = [sample + "::" + neuron for sample, neuron in pairs]
        if frame.NeuronUID.astype(str).tolist() != expected:
            raise ValueError(f"{sheet_name} NeuronUID disagrees with SampleID/NeuronID")
    return tuple(pairs)


def validate_workbook_membership(sheets):
    """Require complete, ordered full identities in every quantitative sheet."""
    required = {f"Projection_{metric}{level}_{side}"
                for metric in ("Length", "Strength")
                for level in ("", "_L3") for side in ("ipsi", "contra")}
    missing = required - set(sheets)
    if missing:
        raise ValueError(f"Workbook missing required quantitative sheets: {sorted(missing)}")
    if "Summary" not in sheets:
        raise ValueError("Workbook missing Summary")
    expected = _ordered_uids(sheets["Summary"], "Summary")
    if not expected:
        raise ValueError("Summary cannot be empty")
    signatures = {"Summary": expected}
    for name, frame in sheets.items():
        if name.startswith("Projection_"):
            actual = _ordered_uids(frame, name)
            if actual != expected:
                raise ValueError(f"{name} ordered membership differs from Summary")
            signatures[name] = actual
    return signatures


def _enrich_henry_layers(workbook):
    layer_script = os.path.join(GROUP_DIR, "scripts", "06a_merge_henry_layers.py")
    spec = importlib.util.spec_from_file_location("merge_henry_layers_phase06a", layer_script)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load layer enrichment: {layer_script}")
    layer_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(layer_module)
    layer_module.merge_henry_layers(workbook)


def harmonize(combined_xlsx=None, out_xlsx=None) -> None:
    """Write a new derivative; existing workbooks are never replaced.

    Both canonical and staging destinations must use a fresh path. This makes
    the historical 306-input/353-destination mismatch an early refusal rather
    than silent membership loss. Curated reference and distant candidate rows
    retain their existing provenance and membership.
    """
    source = Path(combined_xlsx or COMBINED_XLSX).resolve()
    destination = Path(out_xlsx or OUT_XLSX).resolve()
    if os.path.normcase(str(source)) == os.path.normcase(str(destination)):
        raise ValueError("Harmonized input and output must be distinct paths")
    if destination.exists():
        raise FileExistsError(f"Harmonized output already exists; use a new derivative path: {destination}")
    if not source.is_file():
        raise FileNotFoundError(source)

    # Read tabular values for label decisions and complete membership checks.
    # The write below starts from the original workbook, not these dataframes.
    with pd.ExcelFile(source) as xl:
        original_sheet_names = list(xl.sheet_names)
        sheets = {name: pd.read_excel(xl, sheet_name=name) for name in xl.sheet_names}
    input_membership = validate_workbook_membership(sheets)

    summ = sheets["Summary"].copy()
    if "Soma_Region_Refined" not in summ.columns:
        raise KeyError("Summary missing Soma_Region_Refined column")

    # Preserve current state before edit.
    summ["Soma_Region_Refined_PreHarmonize"] = summ["Soma_Region_Refined"]

    is_atlas_only = summ["Soma_Region_Source"] == "auto_atlas_insula"
    print(f"[06] auto_atlas_insula rows to consider: {is_atlas_only.sum()}")

    changed_rows = 0
    ambiguous_rows = 0
    for idx in summ.index[is_atlas_only]:
        atlas_label = summ.at[idx, "Soma_Region_Auto"]
        atlas_leaf = _strip_prefix(atlas_label)
        if atlas_leaf not in EMPIRICAL_MAPPING:
            continue
        manual, conf = EMPIRICAL_MAPPING[atlas_leaf]
        prev = summ.at[idx, "Soma_Region_Refined"]
        if conf == "user_ig":
            if prev != manual:
                summ.at[idx, "Soma_Region_Refined"] = manual
                changed_rows += 1
            new_source = "atlas_Ig_to_IG_granular_user_rule_20260926"
        elif manual == prev:
            # Already in the manual scheme; just update the source flag.
            new_source = "atlas_already_in_manual_scheme"
        else:
            summ.at[idx, "Soma_Region_Refined"] = manual
            changed_rows += 1
            new_source = (
                "atlas_to_manual_harmonized_251637_rule_ambiguous"
                if conf == "ambiguous"
                else "atlas_to_manual_harmonized_251637_rule"
            )
            if conf == "ambiguous":
                ambiguous_rows += 1
        summ.at[idx, "Soma_Region_Source"] = new_source

    print(f"[06] rows with refined label changed: {changed_rows} (ambiguous: {ambiguous_rows})")

    # Per-sample report after harmonization.
    print("[06] Soma_Region_Refined per sample after harmonization:")
    print(
        summ.groupby("SampleID")["Soma_Region_Refined"].value_counts().to_string()
    )
    print("[06] Soma_Region_Source per sample after harmonization:")
    print(
        summ.groupby("SampleID")["Soma_Region_Source"].value_counts().to_string()
    )

    sheets["Summary"] = summ

    # Mapping_Rule sheet documenting the empirical rule.
    def _mapping_notes(leaf: str, conf: str) -> str:
        if conf == "user_ig":
            return (
                "USER RULE 2026-09-26: atlas Ig → granular IG "
                "(overrides empirical Ig→IDD5 crosstab); atlas G untouched"
            )
        if conf == "high":
            rows = REFERENCE_CROSSTAB[REFERENCE_CROSSTAB.atlas_leaf.eq(leaf)]
            dominant = EMPIRICAL_MAPPING[leaf][0]
            fraction = rows.loc[rows.manual_leaf.eq(dominant), "n"].sum() / rows.n.sum()
            return f"{fraction:.1%} concordance in 251637 empirical crosstab"
        return "split assignment in 251637; verify by coord/visual QC"

    mapping_rule = pd.DataFrame(
        [
            {
                "atlas_leaf": k,
                "manual_dominant": v[0],
                "confidence": v[1],
                "source": (
                    "user_rule_20260926_granular_IG"
                    if v[1] == "user_ig"
                    else "empirical_251637_crosstab_n=260"
                ),
                "notes": _mapping_notes(k, v[1]),
            }
            for k, v in EMPIRICAL_MAPPING.items()
        ]
    )
    sheets["Mapping_Rule"] = mapping_rule
    sheets["Mapping_Rule_Crosstab"] = REFERENCE_CROSSTAB

    # Regenerate Provenance counts.
    sheets["Provenance"] = (
        summ["Soma_Region_Source"]
        .value_counts()
        .rename_axis("Soma_Region_Source")
        .reset_index(name="n")
    )

    out_dir = str(destination.parent)
    fd, temp_out = tempfile.mkstemp(
        prefix=".harmonized_build_", suffix=".xlsx", dir=out_dir
    )
    os.close(fd)
    try:
        # Preserve source workbook objects, formulas and formatting. Only the
        # documented Summary fields and generated rule/provenance sheets change.
        workbook = load_workbook(source)
        summary_ws = workbook["Summary"]
        headers = {cell.value: cell.column for cell in summary_ws[1] if cell.value is not None}
        for column in ("Soma_Region_Refined_PreHarmonize", "Soma_Region_Refined", "Soma_Region_Source"):
            if column not in headers:
                headers[column] = summary_ws.max_column + 1
                summary_ws.cell(1, headers[column], column)
            for row_number, value in enumerate(summ[column], start=2):
                summary_ws.cell(row_number, headers[column]).value = None if pd.isna(value) else value
        for name in ("Provenance", "Mapping_Rule", "Mapping_Rule_Crosstab"):
            position = workbook.sheetnames.index(name) if name in workbook.sheetnames else len(workbook.sheetnames)
            if name in workbook:
                del workbook[name]
            sheet = workbook.create_sheet(name, position)
            for row in dataframe_to_rows(sheets[name], index=False, header=True):
                sheet.append([None if pd.isna(value) else value for value in row])
        workbook.save(temp_out)
        workbook.close()

        # The canonical harmonized workbook must include Henry's 251637 layer
        # annotations. Enrich every rebuild so Phase 6 cannot erase Phase 6a.
        _enrich_henry_layers(temp_out)

        # A harmonized build may change labels, never membership.
        built_sheets = pd.read_excel(temp_out, sheet_name=None)
        if validate_workbook_membership(built_sheets) != input_membership:
            raise RuntimeError("Harmonized workbook membership differs from its input")
        if not set(original_sheet_names).issubset(built_sheets):
            raise RuntimeError("Harmonized workbook lost an input sheet")

        # The same-directory hard link publishes a complete file atomically
        # and fails if another process created the destination during the build.
        # os.replace would silently clobber that newly created workbook.
        os.link(temp_out, destination)
    finally:
        if os.path.exists(temp_out):
            os.remove(temp_out)

    print(f"[06] wrote {destination}")
    print("[06] done")


if __name__ == "__main__":
    harmonize()
